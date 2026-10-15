"""Пакет эксперту для агента «Карточка сигнала» (#94, спека #93): справочники и разметка 50 сигналов.

Справочники — версионируемые данные в репозитории (`signal_card_dictionaries.json`): их же
потом читает агент, версия идёт в рецепт карточки. Таблица эксперту — лист справочников на
утверждение и лист разметки: 35 принятых статей ленты за август–октябрь и 15 сигналов радара
(технологии и бизнес), по темам вперемешку.

Разметка вслепую: в таблице только то, что написал издатель, — заголовок, текст, ссылка, дата
публикации. Сути, перевода, темы, решений судьи радара в ней нет. Отсюда правила выборки радара:
главная публикация — статья корпуса или страница, которую докачали (`full_text_fetched`); без
докачки текст находки «режима ChatGPT» — сводка модели. Дата — тоже издателя: у находки «режима
ChatGPT» дата публикации — это дата события по версии модели, её в таблице нет.

Известные пределы данных, а не выгрузки: текст докачанной страницы платформа хранит только
началом (`signal_discovery.WEB_EVIDENCE_FULLTEXT_CHARS`, 1 500 знаков); у страницы без og:title,
twitter:title и h1 докачка оставляет заголовок модели (`parse_article_page`), и по базе его не
отличить от заголовка издателя.

Таблица и манифест (идентификаторы, отпечаток текста) пишутся вне репозитория — он публичный.
Манифест фиксирует входы экзамена: текст публикации агент и экзамен берут через
publication_text — отпечаток совпадёт с тем, что видел эксперт.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from typing import Any

from psycopg.rows import dict_row

from oiltech_digest import config, feed_window
from oiltech_digest.db import repository

DICTIONARIES_PATH = Path(__file__).with_name("signal_card_dictionaries.json")
REPO_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_NAME = "signal_card_markup.xlsx"
MANIFEST_NAME = "signal_card_markup_manifest.json"
MARKUP_SHEET = "Разметка"
DICTIONARY_SHEET = "Справочники"
LISTS_SHEET = "_списки"
UNDEFINED = "не определено"
TEXT_LIMIT = 6000
FEED_COUNT = 35
RADAR_QUOTA = {"technology": 10, "business": 5}
RADAR_COUNT = sum(RADAR_QUOTA.values())
CATEGORY_RU = {"technology": "технология", "business": "бизнес"}
DEFAULT_SEED = 20261014
DEFAULT_SINCE = "2026-08-01"
PERIOD_END = date(2026, 11, 1)  # август–октябрь
# Символы, которые openpyxl в ячейку не пишет (openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE): один такой
# в тексте издателя ронял бы всю выгрузку.
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

# Колонки, которые заполняет эксперт: (заголовок, справочник для выпадающего списка или None).
MARKUP_COLUMNS = [
    ("Дата события", None),
    ("Источник даты события", "date_source"),
    ("Компании-участники", None),
    ("Тип события", "event_type"),
    ("Регион", "region"),
    ("Класс технологии", "technology_class"),
    ("Зрелость", "maturity"),
    ("Связь с нефтесервисом", "oilfield_relevance"),
    ("Применение в нефтесервисе", None),
    ("Почему важно нефтесервису", None),
    ("Комментарий эксперта", None),
]
_GIVEN_COLUMNS = ["ID", "Заголовок", "Текст публикации", "Ссылка", "Дата публикации"]
# На утверждение — четыре справочника спеки и алиасы компаний. Источник даты и связь с
# нефтесервисом — значения самой карточки (спека #93), их эксперт не утверждает, а только ставит.
_APPROVAL = ["event_type", "region", "technology_class", "maturity"]

INSTRUCTION = [
    "Разметка 50 сигналов для проверки агента «Карточка сигнала».",
    "Лист «Справочники»: проверьте значения; в колонке «Решение» — «да», «убрать» или своя правка.",
    "Лист «Разметка»: по каждой строке прочитайте текст публикации и заполните колонки справа.",
    "Размечайте по тексту в таблице — агент увидит только его; ссылка — чтобы проверить контекст.",
    "Дата события — ГГГГ-ММ-ДД, если она названа в тексте или следует из него («из текста»); иначе —",
    "  дата публикации и «по дате публикации»; даты публикации нет — «не определено».",
    "Компании — только участники события, через «;»; издатель публикации участником не считается.",
    "Колонки со справочниками — выпадающий список. «Неприменимо» и «нет события» — атрибут к событию",
    "  не относится (у сделки нет зрелости технологии, у обзора нет события); «не определено» — по",
    "  тексту понять нельзя.",
    "Применение и «почему важно» — одной фразой. Сомнения — в «Комментарий эксперта».",
]


def load_dictionaries() -> dict[str, Any]:
    return json.loads(DICTIONARIES_PATH.read_text(encoding="utf-8"))


def publication_text(raw: str | None) -> str:
    """Текст публикации, который видит эксперт и получает агент: без управляющих символов, до TEXT_LIMIT."""
    return _CONTROL_RE.sub(" ", raw or "").strip()[:TEXT_LIMIT]


def default_until() -> str:
    """Окно закрыто сегодняшним днём (по Москве): то, что ещё собирается, выборку не сдвигает."""
    return min(datetime.now(feed_window.MSK).date(), PERIOD_END).isoformat()


def select_sample(*, seed: int, since: str, until: str) -> list[dict[str, Any]]:
    """Статьи ленты и сигналы радара, темы по кругу. Те же seed и окно — те же сигналы."""
    feed = _round_robin(_feed_candidates(since, until), FEED_COUNT, seed)
    taken = {item["article_id"] for item in feed}
    pool = [item for item in _radar_candidates(until) if item.get("article_id") not in taken]
    radar: list[dict[str, Any]] = []
    for category, quota in RADAR_QUOTA.items():
        radar += _round_robin([item for item in pool if item["category_code"] == category], quota, seed)
    if len(radar) < RADAR_COUNT:  # одной категории не хватило — добираем другой, недобор покажет shortfall
        chosen = {item["id"] for item in radar}
        radar += _round_robin([item for item in pool if item["id"] not in chosen], RADAR_COUNT - len(radar), seed)
    return feed + radar


def shortfall(sample: list[dict[str, Any]]) -> list[str]:
    """Чего не хватает до набора экзамена: 35 + 15 и обе категории на радаре."""
    origin = Counter(item["origin"] for item in sample)
    problems = []
    if origin["feed"] < FEED_COUNT:
        problems.append(f"статей ленты {origin['feed']} из {FEED_COUNT}")
    if origin["radar"] < RADAR_COUNT:
        problems.append(f"сигналов радара {origin['radar']} из {RADAR_COUNT}")
    categories = {item["category_code"] for item in sample if item["origin"] == "radar"}
    problems += [f"на радаре нет категории «{CATEGORY_RU[code]}»" for code in RADAR_QUOTA if code not in categories]
    return problems


def export_pack(out_dir: str | Path | None, *, seed: int | None = None, since: str | None = None,
                until: str | None = None, apply: bool = False) -> dict[str, Any]:
    seed = DEFAULT_SEED if seed is None else seed
    since, until = since or DEFAULT_SINCE, until or default_until()
    sample = select_sample(seed=seed, since=since, until=until)
    result = {"composition": composition(sample), "count": len(sample), "shortfall": shortfall(sample),
              "seed": seed, "since": since, "until": until, "applied": False}
    if not apply:
        return result
    if result["shortfall"]:
        raise PackError("выборка неполная: " + "; ".join(result["shortfall"])
                        + " — экзамен на таком наборе ничего не скажет")
    out = _checked_out_dir(out_dir)
    dictionaries = load_dictionaries()
    out.mkdir(parents=True, exist_ok=True)
    _write_workbook(out / WORKBOOK_NAME, sample, dictionaries)
    manifest = {
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "seed": seed, "since": since, "until": until,
        "dictionaries_version": dictionaries["version"], "text_limit": TEXT_LIMIT,
        "composition": result["composition"],
        "items": [{key: item.get(key) for key in (
            "id", "origin", "article_id", "signal_id", "evidence_id", "url", "category", "theme",
            "text_sha256", "text_chars")} for item in sample],
    }
    (out / MANIFEST_NAME).write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {**result, "applied": True, "workbook": str(out / WORKBOOK_NAME), "manifest": str(out / MANIFEST_NAME)}


def composition(sample: list[dict[str, Any]]) -> dict[str, Any]:
    return {key: dict(sorted(Counter(item[key] for item in sample).items()))
            for key in ("origin", "category", "theme")}


class PackError(ValueError):
    pass


def _checked_out_dir(out_dir: str | Path | None) -> Path:
    if not out_dir:
        raise PackError("нужен --out: папка вне репозитория")
    out = Path(out_dir).expanduser().resolve()
    if out == REPO_ROOT or REPO_ROOT in out.parents:
        raise PackError(f"{out} — внутри репозитория, а он публичный: укажите папку вне его")
    existing = [name for name in (WORKBOOK_NAME, MANIFEST_NAME) if (out / name).exists()]
    if existing:
        raise PackError(f"в {out} уже есть {', '.join(existing)} — не затираю разметку; укажите другую папку")
    return out


def _feed_candidates(since: str, until: str) -> list[dict[str, Any]]:
    """Принятые отбором статьи (relevant = TRUE), видимые в ленте, собранные в окне."""
    with repository.get_connection() as conn:
        cur = conn.cursor(row_factory=dict_row)
        cur.execute(
            f"""
            SELECT a.id AS article_id, a.title, a.url, a.raw_text AS text,
                   COALESCE(a.published_at, a.collected_at) AS published,
                   COALESCE(p.name, t.name, 'без темы') AS theme, 'business' AS category_code
            FROM articles a
            JOIN article_cards c ON c.article_id = a.id
            JOIN sources s ON s.id = a.source_id
            LEFT JOIN article_tags at ON at.article_id = a.id
            LEFT JOIN tags t ON t.id = at.tag_id
            LEFT JOIN tags p ON p.id = t.parent_id
            WHERE c.relevant IS TRUE AND {feed_window.visible_sql('a', 'c', 's')}
              AND a.collected_at >= %s::date AND a.collected_at < %s::date
              AND a.full_text_status IS DISTINCT FROM 'mismatch'
              AND length(btrim(COALESCE(a.raw_text, ''))) >= %s
            ORDER BY a.id
            """,
            (since, until, config.MIN_ARTICLE_TEXT_CHARS),
        )
        rows = cur.fetchall()
    return [_item("feed", f"article:{row['article_id']}", row) for row in rows]


def _radar_candidates(until: str) -> list[dict[str, Any]]:
    """Сигналы радара (технология и бизнес), поступившие до конца окна, с публикацией, текст
    которой написал издатель. Дубли, архив и отсеянное машиной — мимо: в экзамене нужны события,
    а не брак поиска. Правила экрана радара (_radar_quality_sql) здесь не годятся: бизнес-сигналы
    они прячут целиком, а «не про нефтесервис» нужно экзамену — это одно из значений атрибута."""
    from oiltech_digest.signal_research import EVIDENCE_SOURCE

    with repository.get_connection() as conn:
        cur = conn.cursor(row_factory=dict_row)
        cur.execute(
            f"""
            SELECT DISTINCT ON (s.id)
                   s.id AS signal_id, s.signal_category AS category_code, s.theme,
                   e.id AS evidence_id, a.id AS article_id,
                   COALESCE(a.title, e.title) AS title, COALESCE(a.url, e.source_url) AS url,
                   COALESCE(a.raw_text, e.extracted_fact) AS text,
                   CASE WHEN a.id IS NOT NULL THEN COALESCE(a.published_at, a.collected_at)
                        WHEN e.raw_payload_json->>'evidence_source' IS DISTINCT FROM %s THEN e.published_at
                   END AS published
            FROM signals s
            JOIN signal_evidence e ON e.signal_id = s.id
            LEFT JOIN articles a ON a.id = e.article_id
            WHERE s.merged_into_signal_id IS NULL AND s.archived_at IS NULL AND s.filter_stage IS NULL
              -- Только находки поиска (#95): сигнал ленты — это статья, её берёт лентовая часть.
              AND {repository.radar_origin_sql("s")}
              AND s.signal_category IN ('technology', 'business') AND s.first_seen_at < %s::date
              AND (length(btrim(COALESCE(a.raw_text, ''))) >= %s
                   OR (a.id IS NULL AND e.raw_payload_json->>'full_text_fetched' = 'true'
                       AND length(btrim(COALESCE(e.extracted_fact, ''))) >= %s))
            ORDER BY s.id, e.strength DESC, e.id
            """,
            (EVIDENCE_SOURCE, until, config.MIN_ARTICLE_TEXT_CHARS, config.MIN_ARTICLE_TEXT_CHARS),
        )
        rows = cur.fetchall()
    return [_item("radar", f"signal:{row['signal_id']}", row) for row in rows]


def _item(origin: str, item_id: str, row: dict[str, Any]) -> dict[str, Any]:
    text = publication_text(row.pop("text"))
    code = row.pop("category_code")
    return {
        **row, "origin": origin, "id": item_id,
        "title": _CONTROL_RE.sub(" ", row.get("title") or "").strip(), "text": text,
        "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(), "text_chars": len(text),
        "published": _date(row.get("published")), "category_code": code,
        "category": CATEGORY_RU.get(code, code),
    }


def _round_robin(items: list[dict[str, Any]], count: int, seed: int) -> list[dict[str, Any]]:
    """По одному из каждой темы по кругу — чтобы большая тема не съела выборку. Порядок внутри
    темы — по хэшу seed и номера: новый кандидат не перетасовывает остальных, а лишь встаёт на своё место."""
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in items:
        groups[str(item["theme"])].append(item)
    queues = [sorted(group, key=lambda item: hashlib.sha256(f"{seed}:{item['id']}".encode()).hexdigest())
              for _, group in sorted(groups.items())]
    chosen: list[dict[str, Any]] = []
    depth = 0
    while len(chosen) < count and any(depth < len(queue) for queue in queues):
        chosen += [queue[depth] for queue in queues if depth < len(queue)][:count - len(chosen)]
        depth += 1
    return chosen


def _date(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).date().isoformat()
    return str(value)[:10]


def _write_workbook(path: Path, sample: list[dict[str, Any]], dictionaries: dict[str, Any]) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    book = openpyxl.Workbook()
    intro = book.active
    intro.title = "Инструкция"
    for line in INSTRUCTION + ["", f"Версия справочников: {dictionaries['version']}"]:
        intro.append([line])
    intro.column_dimensions["A"].width = 110

    approve = book.create_sheet(DICTIONARY_SHEET)
    approve.append(["Справочник", "Значение", "Правило / алиасы", "Решение", "Комментарий эксперта"])
    for key in _APPROVAL:
        block = dictionaries[key]
        for number, value in enumerate(block["values"]):
            approve.append([block["title"], value, block.get("rule", "") if number == 0 else "", None, None])
    for company in dictionaries["companies"]:
        approve.append(["Компании", company["name"], "; ".join(company["aliases"]), None, None])
    for letter, width in zip("ABCDE", (20, 34, 60, 14, 40)):
        approve.column_dimensions[letter].width = width

    # Значения выпадающих списков — на скрытом листе, а не строкой в самой проверке: у строки
    # предел Excel 255 знаков, а класс технологии с «не определено» длиннее.
    lists = book.create_sheet(LISTS_SHEET)
    lists.sheet_state = "hidden"
    list_refs: dict[str, str] = {}
    for column, key in enumerate(sorted({key for _, key in MARKUP_COLUMNS if key}), start=1):
        values = [*dictionaries[key]["values"], UNDEFINED]
        letter = get_column_letter(column)
        lists.cell(row=1, column=column, value=key)
        for row, value in enumerate(values, start=2):
            lists.cell(row=row, column=column, value=value)
        list_refs[key] = f"'{LISTS_SHEET}'!${letter}$2:${letter}${len(values) + 1}"

    sheet = book.create_sheet(MARKUP_SHEET, index=1)
    sheet.append(_GIVEN_COLUMNS + [name for name, _ in MARKUP_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for item in sample:
        sheet.append([item["id"], item["title"], item["text"], item.get("url") or "", item["published"]])
        for cell in sheet[sheet.max_row]:
            # Заголовок издателя «=…» openpyxl записал бы формулой — в таблице он должен остаться текстом.
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    last = len(sample) + 1
    for offset, (_, key) in enumerate(MARKUP_COLUMNS):
        if not key:
            continue
        letter = get_column_letter(len(_GIVEN_COLUMNS) + offset + 1)
        validation = DataValidation(type="list", formula1=list_refs[key], allow_blank=True)
        validation.add(f"{letter}2:{letter}{last}")
        sheet.add_data_validation(validation)
    widths = [14, 40, 90, 30, 14] + [18] * len(MARKUP_COLUMNS)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "C2"
    book.save(path)

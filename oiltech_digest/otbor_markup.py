"""Загрузка разметки эксперта для «Отбора» (#129, спека #128): проверка целостности и отчёт о нарушениях.

Таблица сверяется с манифестом своего пакета (otbor_pack): номера строк и отпечатки текстов; значения — из тех
списков, что видел эксперт (они в манифесте, справочники с тех пор могли поменяться); правила колонок — те же, что
на листе «Правила»: у строки «не брать» причина отказа, у «брать» тема, общая оценка и баллы по пяти критериям
профиля своей категории с обоснованием. Нарушения собираются все сразу: эксперт правит таблицу за один заход.
Экзамен (#133) берёт отсюда же строки с метками в кодах справочников.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from oiltech_digest import otbor_pack
from oiltech_digest.signal_card import _plain

DECISION_COLUMN = "Брать или нет"
REASON_COLUMN = "Причина отказа"
CATEGORY_COLUMN = "Категория"
# Шаги «Отбора», без которых решения нет, — у каждой строки.
REQUIRED = (DECISION_COLUMN, "Отрасль", "Связь с нефтесервисом", CATEGORY_COLUMN)
# Колонка списка → метка в строке для экзамена.
LABEL_KEYS = {"Отрасль": "industry", "Связь с нефтесервисом": "oilfield_relevance", DECISION_COLUMN: "decision",
              REASON_COLUMN: "reject_reason", CATEGORY_COLUMN: "category", "Тема": "theme",
              otbor_pack.SCALE_COLUMN: "overall"}


class MarkupError(ValueError):
    """Разметку нельзя проверить: нет манифеста, не тот файл, нет листа или колонок."""


def check(markup_path: str | Path, manifest_path: str | Path | None = None) -> dict[str, Any]:
    markup = Path(markup_path).expanduser().resolve()
    manifest_file = (Path(manifest_path).expanduser().resolve() if manifest_path
                     else markup.with_name(otbor_pack.MANIFEST_NAME))
    manifest = _manifest(manifest_file)
    book = _book(markup)
    rows = _markup_rows(book, manifest, markup.name)
    expected = {item["id"]: item for item in manifest["items"]}
    seen: set[str] = set()
    violations: list[str] = []
    items: list[dict[str, Any]] = []
    for cells in rows:
        number = cells["ID"]
        if number not in expected:
            violations.append(f"{number}: строки нет в манифесте")
            continue
        if number in seen:
            violations.append(f"{number}: номер повторяется")
            continue
        seen.add(number)
        problems: list[str] = []
        if _sha(cells.get("Текст публикации")) != expected[number]["text_sha256"]:
            problems.append("текст публикации не тот, что в манифесте")
        items.append({"id": number, "group": expected[number]["group"],
                      **_labels(cells, expected[number], manifest, problems)})
        violations += [f"{number}: {problem}" for problem in problems]
    violations += [f"{number}: строки нет в таблице" for number in expected if number not in seen]
    return {"rows": len(rows), "violations": violations, "items": items, "approvals": _approvals(book),
            "versions": _versions(manifest), "manifest": str(manifest_file)}


def _labels(cells: dict[str, Any], item: dict[str, Any], manifest: dict[str, Any],
            problems: list[str]) -> dict[str, Any]:
    """Метки строки в кодах справочников и нарушения правил колонок."""
    labels: dict[str, Any] = {}
    invalid = set()
    for column, options in manifest["lists"].items():
        text = _text(cells.get(column))
        code = next((one["code"] for one in options if _plain(one["label"]) == _plain(text)), None) if text else None
        if text and code is None:
            problems.append(f"«{column}»: «{text}» нет в списке")
            invalid.add(column)
        labels[LABEL_KEYS.get(column, column)] = code
    problems += [f"«{column}» не заполнена" for column in REQUIRED if not _text(cells.get(column))]
    columns = manifest["score_columns"]
    scores = {one["criterion"]: _score(cells.get(one["header"]), one, problems) for one in columns}
    scored = [one for one in columns if _text(cells.get(one["header"]))]
    decision, category = labels["decision"], labels["category"]
    if decision == "reject":
        if not _text(cells.get(REASON_COLUMN)):
            problems.append(f"«{REASON_COLUMN}» не заполнена")
        if scored:
            problems.append("баллы у строки «не брать»")
    elif decision == "accept":
        if _text(cells.get(REASON_COLUMN)):
            problems.append("у строки «брать» не должно быть причины отказа")
        problems += [f"«{column}» не заполнена" for column in ("Тема", otbor_pack.SCALE_COLUMN)
                     if not _text(cells.get(column)) and column not in invalid]
        profile = otbor_pack.CATEGORY_PROFILE.get(category)
        if category == "other" and scored:
            problems.append("баллы у категории «другое»")
        if profile:
            problems += [f"«{one['header']}» не заполнена" for one in columns
                         if one["profile"] == profile and not _text(cells.get(one["header"]))]
            problems += [f"баллы не той категории: «{one['header']}»" for one in scored if one["profile"] != profile]
            if not _text(cells.get(otbor_pack.WHY_COLUMN)):
                problems.append(f"«{otbor_pack.WHY_COLUMN}» не заполнена")
    if labels["reject_reason"] == "mixed_events" and not item.get("links"):
        problems.append("«Разные события в одной находке» только у строки со ссылками находки")
    return {**labels, "scores": scores, "why": _text(cells.get(otbor_pack.WHY_COLUMN)) or None,
            "comment": _text(cells.get(otbor_pack.COMMENT_COLUMN)) or None}


def _score(raw: Any, column: dict[str, Any], problems: list[str]) -> int | str | None:
    """Балл — целое число от 0 до 100; у «Свежести» ещё «нет даты» (даты события в тексте нет)."""
    text = _text(raw)
    if not text:
        return None
    if column["allows_no_date"] and _plain(text) == _plain(otbor_pack.NO_DATE):
        return otbor_pack.NO_DATE
    try:
        number = float(text.replace(",", "."))
    except ValueError:
        number = -1.0
    if number.is_integer() and 0 <= number <= 100:
        return int(number)
    shown = int(number) if isinstance(raw, (int, float)) and float(raw).is_integer() else text
    problems.append(f"«{column['header']}»: «{shown}» не целое число от 0 до 100"
                    + (f" и не «{otbor_pack.NO_DATE}»" if column["allows_no_date"] else ""))
    return None


def _approvals(book: Any) -> dict[str, Any]:
    """Решения эксперта на листах утверждения: «да», правка или «убрать», без решения."""
    approved, pending, changes, missing = 0, 0, [], []
    for title in (otbor_pack.APPROVAL_SHEET, otbor_pack.RUBRICS_SHEET):
        rows = list(book[title].iter_rows(values_only=True)) if title in book.sheetnames else []
        header = [_text(cell) for cell in rows[0]] if rows else []
        if "Решение" not in header:
            missing.append(f"лист «{title}» с колонкой «Решение»")
            continue
        decision_at = header.index("Решение")
        comment_at = header.index("Комментарий эксперта") if "Комментарий эксперта" in header else None
        for row in rows[1:]:
            if len(row) < 2 or not _text(row[1]):
                continue
            decision = _text(row[decision_at]) if decision_at < len(row) else ""
            comment = _text(row[comment_at]) if comment_at is not None and comment_at < len(row) else ""
            if not decision:
                pending += 1
            elif _plain(decision) == "да":
                approved += 1
            else:
                changes.append(f"{title}: {_text(row[1])}: {decision}" + (f" ({comment})" if comment else ""))
    return {"approved": approved, "pending": pending, "changes": changes, "missing": missing}


def _versions(manifest: dict[str, Any]) -> list[str]:
    """Справочники и линейки, сменившиеся с выгрузки пакета: экзамену нужны те, по которым размечал эксперт."""
    recipe = manifest.get("recipe") or {}
    current = {"dictionaries_version": ("справочники", otbor_pack.load_dictionaries()["version"]),
               "rubrics_version": ("линейки", otbor_pack.load_rubrics()["version"])}
    return [f"{name} пакета {recipe.get(key)}, сейчас {version}" for key, (name, version) in current.items()
            if recipe.get(key) != version]


def _manifest(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise MarkupError(f"нет манифеста пакета {path}: он лежит рядом с таблицей, иначе укажите --manifest")
    data = json.loads(path.read_text(encoding="utf-8"))
    if (not isinstance(data, dict) or data.get("kind") != otbor_pack.MANIFEST_KIND
            or not isinstance(data.get("items"), list)):
        raise MarkupError(f"{path.name}: не манифест пакета «Отбора» (otbor-expert-pack)")
    return data


def _book(path: Path) -> Any:
    import zipfile

    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile) as exc:
        raise MarkupError(f"{path.name}: не таблица xlsx ({exc})") from exc


def _markup_rows(book: Any, manifest: dict[str, Any], name: str) -> list[dict[str, Any]]:
    if otbor_pack.MARKUP_SHEET not in book.sheetnames:
        raise MarkupError(f"{name}: нет листа «{otbor_pack.MARKUP_SHEET}»")
    rows = list(book[otbor_pack.MARKUP_SHEET].iter_rows(values_only=True))
    header = [_text(cell) for cell in rows[0]] if rows else []
    absent = [column for column in ("ID", "Текст публикации", *manifest["columns"]["answers"]) if column not in header]
    if absent:
        raise MarkupError(f"{name}: нет колонок {', '.join(absent)}")
    parsed = []
    for values in rows[1:]:
        cells = dict(zip(header, values))
        number = _text(cells.get("ID"))
        if number:
            parsed.append({**cells, "ID": number})
    return parsed


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _sha(value: Any) -> str:
    return hashlib.sha256(("" if value is None else str(value)).encode("utf-8")).hexdigest()

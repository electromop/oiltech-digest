"""Таблица пакета эксперту «Отбора» (#129): инструкция, разметка, правила колонок, листы на утверждение и скрытые
списки. Состав строк и манифест — otbor_pack. Тексты правил берутся из справочников и линеек, а не пишутся здесь:
одна и та же линейка идёт в правила эксперта и потом в инструкцию модели (спека #128). Тексты для эксперта — без
длинного и среднего тире."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from oiltech_digest.otbor_pack import (APPROVAL_SHEET, CHOICE_COLUMNS, COMMENT_COLUMN, EXPERT_SCALE, GIVEN_COLUMNS,
                                       LISTS_SHEET, MARKUP_SHEET, PROFILE_RU, RUBRICS_SHEET, RULES_SHEET,
                                       SCALE_COLUMN, WHY_COLUMN, answer_columns, score_columns)
from oiltech_digest.scoring_profiles import SCORING_PROFILES

RELEVANCE_COLUMN = "Связь с нефтесервисом"


def write(path: Path, rows: list[dict[str, Any]], *, dictionaries: dict[str, Any], rubrics: dict[str, Any],
          themes: list[dict[str, str]], weights: dict[tuple[str, str], float],
          lists: dict[str, list[dict[str, str]]]) -> None:
    import openpyxl

    book = openpyxl.Workbook()
    intro = book.active
    intro.title = "Инструкция"
    for line in _instruction(dictionaries, rubrics):
        intro.append([line])
    intro.column_dimensions["A"].width = 120
    _markup_sheet(book, rows, rubrics, lists)
    _table(book, RULES_SHEET, _rules(dictionaries, themes), (24, 60, 34, 90))
    _table(book, APPROVAL_SHEET, _approval(dictionaries), (26, 36, 70, 60, 14, 40))
    _table(book, RUBRICS_SHEET, _rubrics(rubrics, weights), (12, 34, 8, 44, 44, 44, 44, 40, 14, 40))
    book.save(path)


def _instruction(dictionaries: dict[str, Any], rubrics: dict[str, Any]) -> list[str]:
    return [
        "Разметка для проверки агента «Отбор» и оценки сигналов по критериям.",
        "«Отбор» решает, брать ли публикацию на платформу: называет отрасль, связь с нефтесервисом, решение, причину "
        "отказа и категорию. Тему и баллы ставит следующий агент, их вы тоже размечаете у взятых строк.",
        "",
        "Порядок работы.",
        f"1. Листы «{APPROVAL_SHEET}» и «{RUBRICS_SHEET}»: проверьте причины отказа, границу прямой связи и линейки "
        "критериев. В колонке «Решение» напишите «да», «убрать» или свою правку.",
        f"2. Лист «{MARKUP_SHEET}»: прочитайте текст публикации и заполните колонки справа по правилам с листа "
        f"«{RULES_SHEET}».",
        "",
        "Размечайте по тексту в таблице: агент получит только его. Ссылка нужна, чтобы проверить контекст.",
        f"У строк, которые вы уже размечали для карточки сигнала, колонка «{RELEVANCE_COLUMN}» заполнена вашим прежним "
        "ответом. Проверьте его по границе прямой связи и при несогласии исправьте.",
        "Баллы ставятся только у строк «брать» с категорией «бизнес» или «технология»: пять колонок этой категории, "
        "от 0 до 100 по линейке. Колонки другой категории оставьте пустыми.",
        f"Сомнения и замечания пишите в «{COMMENT_COLUMN}».",
        "",
        f"Версия справочников: {dictionaries['version']}. Версия линеек: {rubrics['version']}.",
    ]


def _markup_sheet(book: Any, rows: list[dict[str, Any]], rubrics: dict[str, Any],
                  lists: dict[str, list[dict[str, str]]]) -> None:
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    # Значения выпадающих списков — на скрытом листе, а не строкой в самой проверке: у строки предел Excel
    # 255 знаков, а названия тем длиннее.
    hidden = book.create_sheet(LISTS_SHEET)
    hidden.sheet_state = "hidden"
    refs: dict[str, str] = {}
    for number, (column, values) in enumerate(lists.items(), start=1):
        letter = get_column_letter(number)
        hidden.cell(row=1, column=number, value=column)
        for row, value in enumerate(values, start=2):
            hidden.cell(row=row, column=number, value=value["label"])
        refs[column] = f"'{LISTS_SHEET}'!${letter}$2:${letter}${len(values) + 1}"

    sheet = book.create_sheet(MARKUP_SHEET, index=1)
    header = GIVEN_COLUMNS + answer_columns(rubrics)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for item in rows:
        sheet.append([item["id"], item["title"], item["text"], item.get("url") or "", item.get("publisher") or "",
                      item.get("published"), _links_text(item["links"])]
                     + [item.get("relevance") if column == RELEVANCE_COLUMN else None
                        for column in header[len(GIVEN_COLUMNS):]])
        for cell in sheet[sheet.max_row]:
            # Заголовок издателя «=…» openpyxl записал бы формулой: в таблице он должен остаться текстом.
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    last = len(rows) + 1
    letter = {name: get_column_letter(index) for index, name in enumerate(header, start=1)}
    for column in refs:
        validation = DataValidation(type="list", formula1=refs[column], allow_blank=True)
        validation.add(f"{letter[column]}2:{letter[column]}{last}")
        sheet.add_data_validation(validation)
    for one in score_columns(rubrics):
        if one["allows_no_date"]:
            continue  # число или «нет даты» — проверит загрузка разметки
        validation = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=True)
        validation.add(f"{letter[one['header']]}2:{letter[one['header']]}{last}")
        sheet.add_data_validation(validation)
    for name, width in zip(header, [16, 40, 90, 30, 20, 14, 50] + [18] * (len(header) - len(GIVEN_COLUMNS))):
        sheet.column_dimensions[letter[name]].width = width
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=len(GIVEN_COLUMNS)):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "C2"


def _links_text(links: list[dict[str, Any]]) -> str | None:
    if not links:
        return None
    return "\n".join(f"{number}. {link['title']} ({link['publisher']}, {link['published'] or 'дата не указана'})"
                     f"\n{link['url']}" for number, link in enumerate(links, start=1))


def _rules(dictionaries: dict[str, Any], themes: list[dict[str, str]]) -> list[list[Any]]:
    """Правило каждой колонки разметки; значения списков — с правилом и примерами, темы — с описанием."""
    rows: list[list[Any]] = [["Колонка", "Правило", "Значение", "Пояснение и примеры"]]
    for column, key in CHOICE_COLUMNS:
        if key == "theme":
            rows.append([column, "Одна из тем по содержанию публикации. Обязательна у строк «брать».", None, None])
            rows += [[None, None, theme["name"], theme["description"]] for theme in themes]
            continue
        rule = dictionaries[key]["rule"]
        if column == RELEVANCE_COLUMN:
            rule += " У строк прежней разметки карточки колонка заполнена вашим прежним ответом."
        rows.append([column, rule, None, None])
        rows += [[None, None, value["label"], _explained(value)] for value in dictionaries[key]["values"]]
    rows += [
        ["Баллы по критериям", "Только у строк «брать» с категорией «бизнес» или «технология»: пять колонок этой "
         "категории, целое число от 0 до 100 по линейке с листа «Линейки». Колонки другой категории пустые, у "
         "категории «другое» баллов нет. «Свежесть»: «нет даты», если даты события в тексте нет.", None, None],
        [WHY_COLUMN, "Одной-двумя фразами: что в тексте определило баллы. Обязательна, если стоят баллы.", None, None],
        [SCALE_COLUMN, "Ваша оценка строки по вашей шкале. Обязательна у строк «брать».", None, None],
        *[[None, None, short, full.replace(" — ", ": ")] for _, short, full in EXPERT_SCALE],
        [COMMENT_COLUMN, "Свободно: сомнения, замечания к правилам и спискам.", None, None],
        ["Ссылки находки", "Заполнять не нужно: ссылки находки радара, если их несколько. По ним ставится причина "
         "«Разные события в одной находке».", None, None],
    ]
    return rows


def _explained(value: dict[str, Any]) -> str | None:
    parts = [value.get("rule") or ""]
    if value.get("examples"):
        parts.append("Примеры: " + "; ".join(value["examples"]) + ".")
    return " ".join(part for part in parts if part) or None


def _approval(dictionaries: dict[str, Any]) -> list[list[Any]]:
    """На утверждение: шесть причин отказа, граница прямой связи и правило переносимой."""
    rows: list[list[Any]] = [["Раздел", "Значение", "Правило", "Примеры", "Решение", "Комментарий эксперта"]]
    rows += [["Причины отказа", value["label"], value["rule"], "; ".join(value["examples"]), None, None]
             for value in dictionaries["reject_reason"]["values"]]
    sections = {"direct": "Граница прямой связи", "transferable": "Правило переносимой связи"}
    rows += [[sections[value["code"]], value["label"], value["rule"], "; ".join(value["examples"]), None, None]
             for value in dictionaries["oilfield_relevance"]["values"] if value["code"] in sections]
    return rows


def _rubrics(rubrics: dict[str, Any], weights: dict[tuple[str, str], float]) -> list[list[Any]]:
    """Линейки на утверждение: строка — критерий, колонки — что значат ступени балла, с примером."""
    bands = rubrics["bands"]
    rows: list[list[Any]] = [["Профиль", "Критерий", "Вес", *(band["label"] for band in bands), "Примечание",
                              "Решение", "Комментарий эксперта"]]
    for profile in SCORING_PROFILES:
        for criterion in rubrics["profiles"][profile]:
            rows.append([PROFILE_RU[profile], criterion["name"], weights.get((profile, criterion["name"])),
                         *(f"{criterion['bands'][band['code']]['rule']} Пример: "
                           f"{criterion['bands'][band['code']]['example']}." for band in bands),
                         criterion.get("note"), None, None])
    return rows


def _table(book: Any, title: str, rows: list[list[Any]], widths: tuple[int, ...]) -> None:
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    sheet = book.create_sheet(title)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for number, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(number)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"

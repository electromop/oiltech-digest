"""Пакет эксперту для агента «Отбор» (#129, спека #128): справочники, черновики линеек, разметка вслепую.

Справочники «Отбора» (отрасль, связь с нефтесервисом, решение, причины отказа, категория) и черновики линеек
десяти критериев двух профилей — версионируемые данные в репозитории (`otbor_dictionaries.json`,
`scoring_rubrics.json`): их же потом читают агенты (#130 — «Отбор», #132 — второй вызов «Карточки»), версии
идут в рецепт. Причины отказа, границу прямой связи и линейки эксперт утверждает в том же пакете.

Состав — решения владельца 21-22.10 (спека #128, «Набор и экзамен»), около 150-160 строк:
- прежняя разметка карточки (#94): связь с нефтесервисом у этих строк уже размечена — колонка заполнена
  ответом самого эксперта; входы — из прежнего манифеста (там поправки с первоисточника), а не из базы;
- случайные статьи ленты за 30 дней до --until: 50 отклонённых гейтом и 30 принятых; окно — в манифесте;
- 15-20 находок из «Отсеянных и скрытых» радара: брак судьи и ревью пачки, смешанные ссылки, старые события
  (по очереди из каждого вида);
- находки радара с несколькими ссылками — для причины «разные события в одной находке»;
- накопленные материалы коллег, ещё не ставшие сигналом.

Вслепую: в таблице то, что получит агент, — заголовок, текст публикации, ссылка, издатель, дата публикации,
у находки радара — список её ссылок. Решений гейта, судьи и поиска, сути, темы, оценки автора материала в
ней нет; группа строки не видна ни колонкой, ни порядком. Правила текста и издателя — как у пакета карточки
(signal_card_pack): у находки «режима ChatGPT» дата публикации — дата события по версии модели поиска, а имя
издателя записала модель, поэтому в таблице сайт из ссылки и пустая дата.

Таблица и манифест пишутся вне репозитория — он публичный. Проверка размеченной таблицы — otbor_markup.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
from typing import Any

from psycopg.rows import dict_row

from oiltech_digest import colleague_materials, config, feed_window
from oiltech_digest.db import repository
from oiltech_digest.scoring_profiles import SCORING_PROFILES
from oiltech_digest.signal_card_pack import (REPO_ROOT, TEXT_LIMIT, _CONTROL_RE, _file_mark, _markable,
                                             _publisher_label, _rank, publication_day, publication_text)

DICTIONARIES_PATH = Path(__file__).with_name("otbor_dictionaries.json")
RUBRICS_PATH = Path(__file__).with_name("scoring_rubrics.json")
WORKBOOK_NAME = "otbor_markup.xlsx"
MANIFEST_NAME = "otbor_markup_manifest.json"
MANIFEST_KIND = "otbor_expert_pack"
MARKUP_SHEET = "Разметка"
RULES_SHEET = "Правила"
APPROVAL_SHEET = "На утверждение"
RUBRICS_SHEET = "Линейки"
LISTS_SHEET = "_списки"
GROUPS = {
    "card_markup": "прежняя разметка карточки",
    "feed_rejected": "отклонённые гейтом статьи ленты",
    "feed_accepted": "принятые гейтом статьи ленты",
    "radar_filtered": "отсеянные и скрытые находки радара",
    "radar_links": "находки с несколькими ссылками",
    "materials": "материалы коллег",
}
QUOTAS = {"feed_rejected": 50, "feed_accepted": 30, "radar_filtered": 18, "radar_links": 10}
# Меньше этого экзамену группы не хватит: «Отсеянных» — 15-20 (спека), находок со ссылками — хотя бы пять.
MINIMUMS = {"feed_rejected": 50, "feed_accepted": 30, "radar_filtered": 15, "radar_links": 5}
FILTERED_KINDS = {"judge": "брак судьи", "review": "брак ревью пачки", "mixed": "смешанные ссылки",
                  "old": "старые события"}
FEED_WINDOW_DAYS = 30
LINKS_LIMIT = 8  # ссылок находки во входе «Отбора» (спека #128)
DEFAULT_SEED = 20261022
NO_DATE = "нет даты"
PROFILE_RU = {"business": "Бизнес", "tech_radar": "Технология"}
# Профиль баллов — по категории (спека #128): «другое» баллами не оценивается.
CATEGORY_PROFILE = {"business": "business", "technology": "tech_radar"}
GIVEN_COLUMNS = ["ID", "Заголовок", "Текст публикации", "Ссылка", "Издатель", "Дата публикации", "Ссылки находки"]
# Колонки «Отбора» — в порядке его рассуждения (спека #128): отрасль, связь, решение, причина, категория. Тема —
# от «Карточки», но размечается здесь же: у взятых строк по ней мерится основной вызов «Карточки».
CHOICE_COLUMNS = [("Отрасль", "industry"), ("Связь с нефтесервисом", "oilfield_relevance"),
                  ("Брать или нет", "decision"), ("Причина отказа", "reject_reason"), ("Категория", "category"),
                  ("Тема", "theme")]
WHY_COLUMN = "Почему такой балл"
SCALE_COLUMN = "Общая оценка"
COMMENT_COLUMN = "Комментарий эксперта"
# Общая оценка строки — шкала эксперта (его список от 13.09, тот же, что у материалов коллег и карточки радара):
# пункты-оценки. «Не тот блок», «дубль» и «оценка завышена» — действия над сигналом, а не оценка строки.
EXPERT_SCALE = [one for one in colleague_materials.VERDICTS
                if one[0] not in ("wrong_block", "merge_duplicate", "overrated")]
_MANIFEST_KEYS = ("id", "group", "kind", "origin", "article_id", "signal_id", "evidence_id", "material_id", "url",
                  "title", "text", "text_sha256", "text_chars", "published", "publisher", "publisher_type",
                  "language", "links", "relevance", "overrides", "override_note", "file_note")


class PackError(ValueError):
    pass


def load_dictionaries() -> dict[str, Any]:
    return json.loads(DICTIONARIES_PATH.read_text(encoding="utf-8"))


def load_rubrics() -> dict[str, Any]:
    return json.loads(RUBRICS_PATH.read_text(encoding="utf-8"))


def score_columns(rubrics: dict[str, Any]) -> list[dict[str, Any]]:
    """Колонки баллов: пять критериев бизнеса и пять технологии; эксперт заполняет колонки своей категории."""
    return [{"header": f"{PROFILE_RU[profile]}: {criterion['name']}", "profile": profile,
             "criterion": criterion["name"], "allows_no_date": bool(criterion.get("allows_no_date"))}
            for profile in SCORING_PROFILES for criterion in rubrics["profiles"][profile]]


def answer_columns(rubrics: dict[str, Any]) -> list[str]:
    return ([name for name, _ in CHOICE_COLUMNS] + [one["header"] for one in score_columns(rubrics)]
            + [WHY_COLUMN, SCALE_COLUMN, COMMENT_COLUMN])


def default_until() -> str:
    """Окно закрыто сегодняшним днём (по Москве): то, что ещё собирается, выборку не сдвигает."""
    return datetime.now(feed_window.MSK).date().isoformat()


def export_pack(out_dir: str | Path | None, *, seed: int | None = None, until: str | None = None,
                apply: bool = False, card_manifest: str | Path | None = None,
                card_markup: str | Path | None = None) -> dict[str, Any]:
    seed = DEFAULT_SEED if seed is None else seed
    until = until or default_until()
    since = (date.fromisoformat(until) - timedelta(days=FEED_WINDOW_DAYS)).isoformat()
    card, card_pack = None, None
    if card_manifest or card_markup:
        card = _card_rows(card_manifest, card_markup)
        card_pack = {**_file_mark("manifest", str(card_manifest)), **_file_mark("markup", str(card_markup))}
    rows = select_rows(seed=seed, since=since, until=until, card=card or [])
    dictionaries, rubrics = load_dictionaries(), load_rubrics()
    weights, blockers = _criteria(rubrics)
    themes = _themes()
    if not themes:
        blockers.append("в базе нет тем (корневых тегов): в колонке «Тема» не из чего выбирать")
    missing = shortfall(rows, card is not None)
    result = {"count": len(rows), "composition": composition(rows), "shortfall": missing, "blockers": blockers,
              "seed": seed, "since": since, "until": until, "applied": False}
    if not apply:
        return result
    if missing or blockers:
        raise PackError("пакет не записан: " + "; ".join(missing + blockers))
    out = _checked_out_dir(out_dir)
    from oiltech_digest import otbor_pack_workbook

    lists = _lists(dictionaries, themes)
    # Строки вперемешку и своим ключом: блоками групп таблица выдавала бы эксперту решение гейта и радара.
    rows = sorted(rows, key=lambda item: _rank("table", seed, item["id"]))
    out.mkdir(parents=True, exist_ok=True)
    otbor_pack_workbook.write(out / WORKBOOK_NAME, rows, dictionaries=dictionaries, rubrics=rubrics, themes=themes,
                              weights=weights, lists=lists)
    manifest = {
        "kind": MANIFEST_KIND, "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "seed": seed, "window": {"since": since, "until": until, "days": FEED_WINDOW_DAYS},
        "recipe": {"dictionaries_version": dictionaries["version"], "rubrics_version": rubrics["version"],
                   "text_limit": TEXT_LIMIT, "links_limit": LINKS_LIMIT},
        "card_pack": card_pack,
        "composition": result["composition"],
        "columns": {"given": GIVEN_COLUMNS, "answers": answer_columns(rubrics)},
        "lists": lists, "score_columns": score_columns(rubrics),
        "items": [{key: item.get(key) for key in _MANIFEST_KEYS} for item in rows],
    }
    (out / MANIFEST_NAME).write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {**result, "applied": True, "workbook": str(out / WORKBOOK_NAME), "manifest": str(out / MANIFEST_NAME)}


def select_rows(*, seed: int, since: str, until: str, card: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Строки пакета по группам. Прежние строки карточки в другие группы не идут. Те же зерно и окно — те же строки:
    порядок в группе — по хэшу зерна, группы и номера, новый кандидат не перетасовывает остальных."""
    rows = list(card)
    articles = {item["article_id"] for item in rows if item.get("article_id")}
    signals = {item["signal_id"] for item in rows if item.get("signal_id")}
    feed = [item for item in _feed_candidates(since, until) if item["article_id"] not in articles and _markable(item)]
    for group in ("feed_rejected", "feed_accepted"):
        pool = [item for item in feed if item["group"] == group]
        rows += sorted(pool, key=lambda item: _rank(seed, group, item["id"]))[:QUOTAS[group]]
    radar = [item for item in _radar_candidates(until) if item["signal_id"] not in signals and _markable(item)]
    rows += _by_kinds([item for item in radar if item["kind"] in FILTERED_KINDS], QUOTAS["radar_filtered"], seed)
    with_links = [item for item in radar if item["kind"] is None and item["link_count"] >= 2]
    rows += sorted(with_links, key=lambda item: _rank(seed, "radar_links", item["id"]))[:QUOTAS["radar_links"]]
    rows += [item for item in _material_candidates(until) if _markable(item)]
    _complete_from_base(rows)
    return rows


def composition(rows: list[dict[str, Any]]) -> dict[str, Any]:
    return {"group": {group: count for group in GROUPS if (count := sum(item["group"] == group for item in rows))},
            "kind": dict(sorted(Counter(item["kind"] for item in rows if item.get("kind")).items()))}


def shortfall(rows: list[dict[str, Any]], card_given: bool) -> list[str]:
    counts = Counter(item["group"] for item in rows)
    missing = [] if card_given else [f"{GROUPS['card_markup']}: нужны --card-manifest и --card-markup"]
    return missing + [f"{GROUPS[group]} {counts[group]} из {minimum}"
                      for group, minimum in MINIMUMS.items() if counts[group] < minimum]


def _checked_out_dir(out_dir: str | Path | None) -> Path:
    if not out_dir:
        raise PackError("нужен --out: папка вне репозитория")
    out = Path(out_dir).expanduser().resolve()
    if out == REPO_ROOT or REPO_ROOT in out.parents:
        raise PackError(f"{out} внутри репозитория, а он публичный: укажите папку вне его")
    existing = [name for name in (WORKBOOK_NAME, MANIFEST_NAME) if (out / name).exists()]
    if existing:
        raise PackError(f"в {out} уже есть {', '.join(existing)}: разметку не затираю, укажите другую папку")
    return out


# --- Группы строк --------------------------------------------------------------------------------


def _card_rows(manifest_path: str | Path | None, markup_path: str | Path | None) -> list[dict[str, Any]]:
    """Прежняя разметка карточки (#94): строки и входы — из её манифеста, связь с нефтесервисом — ответ эксперта
    из заполненной таблицы. Таблица и манифест должны быть об одних и тех же текстах."""
    from oiltech_digest.signal_card import _plain

    if not (manifest_path and markup_path):
        raise PackError(f"{GROUPS['card_markup']}: нужны оба файла, --card-manifest и --card-markup")
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
    items = manifest.get("items") if isinstance(manifest, dict) else None
    if not isinstance(items, list) or not all(
            isinstance(item, dict) and {"id", "text", "text_sha256"} <= item.keys() for item in items):
        raise PackError(f"{Path(manifest_path).name}: не манифест пакета карточки, нужен список items с id, text "
                        "и text_sha256")
    answers = _card_answers(Path(markup_path))
    allowed = {_plain(value["label"]): value["label"] for value in load_dictionaries()["oilfield_relevance"]["values"]}
    rows, problems = [], []
    for item in items:
        answer = answers.get(item["id"])
        if answer is None:
            problems.append(f"{item['id']}: нет строки в таблице")
        elif _sha(item["text"]) != item["text_sha256"] or answer["text_sha256"] != item["text_sha256"]:
            problems.append(f"{item['id']}: текст не тот, что в манифесте")
        elif _plain(answer["relevance"]) not in allowed:
            problems.append(f"{item['id']}: связь «{answer['relevance']}» не из справочника")
        else:
            # Издатель — как его видел эксперт: источник ленты или корпуса, у находки поиска — сайт из ссылки.
            publisher = item.get("publisher") if item.get("publisher_type") else _publisher_label(item)
            rows.append({**{key: item.get(key) for key in ("origin", "article_id", "signal_id", "evidence_id", "url",
                                                            "published", "publisher_type", "override_note")},
                         "id": item["id"], "group": "card_markup", "title": item.get("title") or "",
                         "text": item["text"], "text_sha256": item["text_sha256"], "text_chars": len(item["text"]),
                         "publisher": publisher or "", "relevance": allowed[_plain(answer["relevance"])],
                         "overrides": item.get("overrides") or []})
    extra = sorted(set(answers) - {item["id"] for item in items})
    if extra:
        problems.append("в таблице лишние строки " + ", ".join(extra))
    if problems:
        raise PackError(f"{GROUPS['card_markup']} не сходится с манифестом: " + "; ".join(problems))
    return rows


def _card_answers(path: Path) -> dict[str, dict[str, str]]:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if MARKUP_SHEET not in book.sheetnames:
        raise PackError(f"{path.name}: нет листа «{MARKUP_SHEET}»")
    rows = list(book[MARKUP_SHEET].iter_rows(values_only=True))
    header = [str(cell or "").strip() for cell in rows[0]] if rows else []
    absent = [name for name in ("ID", "Текст публикации", "Связь с нефтесервисом") if name not in header]
    if absent:
        raise PackError(f"{path.name}: нет колонок {', '.join(absent)}")
    answers = {}
    for values in rows[1:]:
        cells = dict(zip(header, values))
        number = str(cells.get("ID") or "").strip()
        if number:
            answers[number] = {"text_sha256": _sha(str(cells.get("Текст публикации") or "")),
                               "relevance": str(cells.get("Связь с нефтесервисом") or "").strip()}
    return answers


def _feed_candidates(since: str, until: str) -> list[dict[str, Any]]:
    """Статьи ленты, собранные в окне, о которых гейт уже решил. Условия ленты (feed_window.visible_sql) — кроме
    решения гейта: отклонённые лента не показывает, а пакету они нужны. Отметки людей выборку не трогают."""
    with repository.get_connection() as conn:
        cur = conn.cursor(row_factory=dict_row)
        cur.execute(
            """
            SELECT a.id AS article_id, a.title, a.url, a.raw_text AS text, a.language,
                   COALESCE(a.published_at, a.collected_at) AS published,
                   s.name AS publisher, COALESCE(s.source_type, '') AS publisher_type, c.relevant
            FROM articles a
            JOIN article_cards c ON c.article_id = a.id
            JOIN sources s ON s.id = a.source_id
            WHERE c.relevant IS NOT NULL AND NOT a.pending_deletion AND s.archived_at IS NULL
              AND NOT EXISTS (SELECT 1 FROM article_reprints ar WHERE ar.article_id = a.id)
              AND a.collected_at >= %s::date AND a.collected_at < %s::date
              AND a.full_text_status IS DISTINCT FROM 'mismatch'
              AND length(btrim(COALESCE(a.raw_text, ''))) >= %s
            ORDER BY a.id
            """,
            (since, until, config.MIN_ARTICLE_TEXT_CHARS),
        )
        rows = cur.fetchall()
    return [_item("feed_accepted" if row.pop("relevant") else "feed_rejected", "feed", f"article:{row['article_id']}",
                  row) for row in rows]


def _radar_candidates(until: str) -> list[dict[str, Any]]:
    """Находки поиска, поступившие до конца окна, с главной публикацией, текст которой написал издатель (как у
    пакета карточки). Вид отсева: брак судьи, брак ревью пачки, смешанные ссылки, событие старше срока на день
    находки (дата судьи или человека, не агента карточки — как у правил показа); прочие этапы отсева — мимо."""
    from oiltech_digest.signal_research import EVIDENCE_SOURCE

    age = int(config.SIGNAL_RADAR_MAX_EVENT_AGE_DAYS)
    old = (f"WHEN {repository._judged_sql('s', 'event_date')} < s.first_seen_at::date - {age} THEN 'old'"
           if age > 0 else "")
    with repository.get_connection() as conn:
        cur = conn.cursor(row_factory=dict_row)
        cur.execute(
            f"""
            SELECT DISTINCT ON (s.id)
                   s.id AS signal_id, e.id AS evidence_id, a.id AS article_id, a.language,
                   COALESCE(a.title, e.title) AS title, COALESCE(a.url, e.source_url) AS url,
                   COALESCE(a.raw_text, e.extracted_fact) AS text,
                   CASE WHEN a.id IS NOT NULL THEN COALESCE(a.published_at, a.collected_at)
                        WHEN e.raw_payload_json->>'evidence_source' IS DISTINCT FROM %s THEN e.published_at
                   END AS published,
                   src.name AS publisher_source, COALESCE(src.source_type, '') AS publisher_type,
                   CASE WHEN s.filter_stage = 'judge' THEN 'judge'
                        WHEN s.filter_stage = 'review' THEN 'review'
                        WHEN s.filter_stage IS NOT NULL THEN 'other_stage'
                        WHEN s.mixed_events IS TRUE THEN 'mixed'
                        {old}
                   END AS kind,
                   (SELECT count(*) FROM signal_evidence x WHERE x.signal_id = s.id) AS link_count
            FROM signals s
            JOIN signal_evidence e ON e.signal_id = s.id
            LEFT JOIN articles a ON a.id = e.article_id
            LEFT JOIN sources src ON src.id = a.source_id
            WHERE {repository.radar_origin_sql("s")} AND s.merged_into_signal_id IS NULL AND s.archived_at IS NULL
              AND s.first_seen_at < %s::date
              AND (length(btrim(COALESCE(a.raw_text, ''))) >= %s
                   OR (a.id IS NULL AND e.raw_payload_json->>'full_text_fetched' = 'true'
                       AND length(btrim(COALESCE(e.extracted_fact, ''))) >= %s))
            ORDER BY s.id, e.strength DESC, e.id
            """,
            (EVIDENCE_SOURCE, until, config.MIN_ARTICLE_TEXT_CHARS, config.MIN_ARTICLE_TEXT_CHARS),
        )
        rows = cur.fetchall()
    items = []
    for row in rows:
        group = "radar_filtered" if row["kind"] in FILTERED_KINDS else "radar_links"
        items.append(_item(group, "radar", f"signal:{row['signal_id']}", {**row, "publisher": _publisher_label(row)}))
    return items


def _material_candidates(until: str) -> list[dict[str, Any]]:
    """Материалы коллег со статусом «принят» (сигналом ещё не стали), присланные до конца окна."""
    with repository.get_connection() as conn:
        cur = conn.cursor(row_factory=dict_row)
        cur.execute("SELECT id, title, description, url, file_path, created_at FROM colleague_materials"
                    " WHERE signal_id IS NULL AND created_at < %s::date ORDER BY id", (until,))
        rows = cur.fetchall()
    items = []
    for row in rows:
        text, note = _material_text(row)
        items.append(_item("materials", "material", f"material:{row['id']}", {
            "material_id": row["id"], "title": row["title"], "text": text, "url": row["url"],
            "published": row["created_at"], "publisher": _publisher_label(row) if row["url"] else "",
            "publisher_type": "", "language": None, "file_note": note}))
    return items


def _material_text(row: dict[str, Any]) -> tuple[str, str | None]:
    """Текст материала — описание автора; у файла и текст файла: так материал пойдёт во вход «Отбора» (спека #128,
    в тех же пределах, что статья). Файл не прочитался — только описание, причина — в манифесте."""
    if not row.get("file_path"):
        return row["description"], None
    from oiltech_digest.documents.model import DocumentError
    from oiltech_digest.documents.parsing import parse_document

    try:
        document = parse_document(Path(row["file_path"]).read_bytes())
    except (OSError, DocumentError) as exc:
        return row["description"], f"файл не прочитан: {exc}"
    return row["description"] + "\n\n" + "\n".join(anchor.text for anchor in document.anchors), None


def _item(group: str, origin: str, item_id: str, row: dict[str, Any]) -> dict[str, Any]:
    text = publication_text(row.get("text"))
    return {**row, "group": group, "kind": row.get("kind") if group == "radar_filtered" else None,
            "origin": origin, "id": item_id, "title": _CONTROL_RE.sub(" ", row.get("title") or "").strip(),
            "text": text, "text_sha256": _sha(text), "text_chars": len(text),
            "published": publication_day(row.get("published")), "links": [], "overrides": []}


def _by_kinds(items: list[dict[str, Any]], count: int, seed: int) -> list[dict[str, Any]]:
    """По одной из каждого вида отсева по кругу — чтобы брак судьи не съел группу; порядок внутри вида — по хэшу."""
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in items:
        groups[item["kind"]].append(item)
    queues = [sorted(group, key=lambda item: _rank(seed, "radar_filtered", item["id"]))
              for _, group in sorted(groups.items())]
    chosen: list[dict[str, Any]] = []
    depth = 0
    while len(chosen) < count and any(depth < len(queue) for queue in queues):
        chosen += [queue[depth] for queue in queues if depth < len(queue)][:count - len(chosen)]
        depth += 1
    return chosen


def _complete_from_base(rows: list[dict[str, Any]]) -> None:
    """Ссылки находки (до LINKS_LIMIT, если их больше одной) — у всех строк радара, в том числе прежних; язык — у
    прежних строк со статьёй (в прежнем манифесте его нет). Правила заголовка, издателя и даты ссылки — как у
    главной публикации: у ссылки поиска сайт из адреса и пустая дата."""
    from oiltech_digest.signal_research import EVIDENCE_SOURCE

    signals = sorted({int(item["signal_id"]) for item in rows if item.get("origin") == "radar" and item.get("signal_id")})
    articles = sorted({int(item["article_id"]) for item in rows if item.get("article_id") and not item.get("language")})
    with repository.get_connection() as conn:
        cur = conn.cursor(row_factory=dict_row)
        cur.execute(
            """
            SELECT e.signal_id, COALESCE(a.title, e.title) AS title, COALESCE(a.url, e.source_url) AS url,
                   src.name AS publisher_source,
                   CASE WHEN a.id IS NOT NULL THEN COALESCE(a.published_at, a.collected_at)
                        WHEN e.raw_payload_json->>'evidence_source' IS DISTINCT FROM %s THEN e.published_at
                   END AS published
            FROM signal_evidence e
            LEFT JOIN articles a ON a.id = e.article_id
            LEFT JOIN sources src ON src.id = a.source_id
            WHERE e.signal_id = ANY(%s::bigint[])
            ORDER BY e.signal_id, e.strength DESC, e.id
            """,
            (EVIDENCE_SOURCE, signals),
        )
        links: dict[int, list[dict[str, Any]]] = defaultdict(list)
        for row in cur.fetchall():
            links[int(row["signal_id"])].append({
                "url": row["url"], "title": _CONTROL_RE.sub(" ", row["title"] or "").strip(),
                "publisher": _publisher_label(row), "published": publication_day(row["published"])})
        cur.execute("SELECT id, language FROM articles WHERE id = ANY(%s::bigint[])", (articles,))
        languages = {int(row["id"]): row["language"] for row in cur.fetchall()}
    for item in rows:
        found = links.get(int(item["signal_id"])) if item.get("origin") == "radar" and item.get("signal_id") else None
        item["links"] = found[:LINKS_LIMIT] if found and len(found) > 1 else []
        if item.get("article_id") and not item.get("language"):
            item["language"] = languages.get(int(item["article_id"]))


# --- Справочная часть ----------------------------------------------------------------------------


def _criteria(rubrics: dict[str, Any]) -> tuple[dict[tuple[str, str], float], list[str]]:
    """Веса критериев с экрана «Скоринг» и расхождения его наборов с линейками: баллы эксперта должны быть о тех
    же критериях, по которым платформа посчитает итог."""
    weights: dict[tuple[str, str], float] = {}
    problems = []
    for profile in SCORING_PROFILES:
        enabled = repository.list_enabled_scoring_criteria(profile)
        screen = [str(one["name"]) for one in enabled]
        drafted = [one["name"] for one in rubrics["profiles"][profile]]
        missing = [name for name in drafted if name not in screen]
        extra = [name for name in screen if name not in drafted]
        if missing or extra:
            problems.append(f"критерии «{PROFILE_RU[profile]}» на экране «Скоринг» не те, что в линейках: "
                            + "; ".join(part for part in (
                                "в линейках есть, на экране нет " + ", ".join(f"«{one}»" for one in missing)
                                if missing else "",
                                "на экране есть, в линейках нет " + ", ".join(f"«{one}»" for one in extra)
                                if extra else "") if part))
        weights.update({(profile, str(one["name"])): _number(one["weight"]) for one in enabled})
    return weights, problems


def _themes() -> list[dict[str, str]]:
    """Темы — корневые тематики заказчика без служебного приёмника, как у радара (topics_from_tags)."""
    from oiltech_digest.signal_discovery import topics_from_tags

    return [{"name": topic["name"], "description": topic["description"]}
            for topic in topics_from_tags(repository.list_enabled_tags())]


def _lists(dictionaries: dict[str, Any], themes: list[dict[str, str]]) -> dict[str, list[dict[str, str]]]:
    """Значения выпадающих списков по колонкам: метка — для эксперта, код — для экзамена. Хранятся в манифесте:
    разметка сверяется с тем, что видел эксперт, даже если справочники потом поменяют."""
    lists = {column: [{"code": value["code"], "label": value["label"]} for value in dictionaries[key]["values"]]
             for column, key in CHOICE_COLUMNS if key != "theme"}
    lists["Тема"] = [{"code": theme["name"], "label": theme["name"]} for theme in themes]
    lists[SCALE_COLUMN] = [{"code": value, "label": short} for value, short, _ in EXPERT_SCALE]
    return lists


def _number(value: Any) -> float | int:
    number = float(value or 0)
    return int(number) if number.is_integer() else number


def _sha(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()

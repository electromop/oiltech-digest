"""Формат пакета эксперту «Отбора» (#129, спека #128): справочники и линейки, листы и колонки таблицы, списки,
отпечаток текста и чтение листа.

Справочники «Отбора» (`otbor_dictionaries.json`) и черновики линеек десяти критериев (`scoring_rubrics.json`) —
версионируемые данные: их же потом читают агенты (#130 — «Отбор», #132 — второй вызов «Карточки»), версии идут в
рецепт. Модуль общий для выгрузки пакета (otbor_pack, otbor_pack_workbook) и проверки разметки (otbor_markup).
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from oiltech_digest import colleague_materials
from oiltech_digest.scoring_profiles import SCORING_PROFILES

DICTIONARIES_PATH = Path(__file__).with_name("otbor_dictionaries.json")
RUBRICS_PATH = Path(__file__).with_name("scoring_rubrics.json")
WORKBOOK_NAME = "otbor_markup.xlsx"
MANIFEST_NAME = "otbor_markup_manifest.json"
MANIFEST_KIND = "otbor_expert_pack"
MARKUP_SHEET = "Разметка"
RULES_SHEET = "Правила"
APPROVAL_SHEET = "На утверждение"
RUBRICS_SHEET = "Линейки"
LISTS_SHEET = "_списки"
NO_DATE = "нет даты"
PROFILE_RU = {"business": "Бизнес", "tech_radar": "Технология"}
# Профиль баллов — по категории (спека #128): «другое» баллами не оценивается.
CATEGORY_PROFILE = {"business": "business", "technology": "tech_radar"}

INDUSTRY_COLUMN = "Отрасль"
RELEVANCE_COLUMN = "Связь с нефтесервисом"
DECISION_COLUMN = "Брать или нет"
REASON_COLUMN = "Причина отказа"
CATEGORY_COLUMN = "Категория"
THEME_COLUMN = "Тема"
WHY_COLUMN = "Почему такой балл"
SCALE_COLUMN = "Общая оценка"
COMMENT_COLUMN = "Комментарий эксперта"
# «Дата поступления» — когда публикация пришла на платформу: от неё отсчитывается старое событие, если издатель
# даты публикации не дал (у находки «режима ChatGPT» её дала модель поиска, в таблицу она не идёт).
GIVEN_COLUMNS = ["ID", "Заголовок", "Текст публикации", "Ссылка", "Издатель", "Дата публикации", "Дата поступления",
                 "Ссылки находки"]
# Колонки «Отбора» — в порядке его рассуждения (спека #128): отрасль, связь, решение, причина, категория. Тема —
# от «Карточки», но размечается здесь же: у взятых строк по ней мерится основной вызов «Карточки». Ключ — справочник.
CHOICE_COLUMNS = [(INDUSTRY_COLUMN, "industry"), (RELEVANCE_COLUMN, "oilfield_relevance"),
                  (DECISION_COLUMN, "decision"), (REASON_COLUMN, "reject_reason"), (CATEGORY_COLUMN, "category"),
                  (THEME_COLUMN, "theme")]
# Общая оценка строки — шкала эксперта (его список от 13.09, тот же, что у материалов коллег и карточки радара):
# пункты-оценки. «Не тот блок», «дубль» и «оценка завышена» — действия над сигналом, а не оценка строки.
EXPERT_SCALE = [one for one in colleague_materials.VERDICTS
                if one[0] not in ("wrong_block", "merge_duplicate", "overrated")]


class PackError(ValueError):
    """Пакет не собрать или его файлы не те: текст — для человека."""


def load_dictionaries() -> dict[str, Any]:
    return json.loads(DICTIONARIES_PATH.read_text(encoding="utf-8"))


def load_rubrics() -> dict[str, Any]:
    return json.loads(RUBRICS_PATH.read_text(encoding="utf-8"))


def score_columns(rubrics: dict[str, Any]) -> list[dict[str, Any]]:
    """Колонки баллов: пять критериев бизнеса и пять технологии; эксперт заполняет колонки своей категории."""
    return [{"header": f"{PROFILE_RU[profile]}: {criterion['name']}", "profile": profile,
             "criterion": criterion["name"], "allows_no_date": bool(criterion.get("allows_no_date"))}
            for profile in SCORING_PROFILES for criterion in rubrics["profiles"][profile]]


def answer_columns(rubrics: dict[str, Any]) -> list[str]:
    return ([name for name, _ in CHOICE_COLUMNS] + [one["header"] for one in score_columns(rubrics)]
            + [WHY_COLUMN, SCALE_COLUMN, COMMENT_COLUMN])


def lists(dictionaries: dict[str, Any], themes: list[dict[str, str]]) -> dict[str, list[dict[str, str]]]:
    """Значения выпадающих списков по колонкам: метка — для эксперта, код — для экзамена. Хранятся в манифесте:
    разметка сверяется с тем, что видел эксперт, даже если справочники потом поменяют."""
    chosen = {column: [{"code": value["code"], "label": value["label"]} for value in dictionaries[key]["values"]]
              for column, key in CHOICE_COLUMNS if key != "theme"}
    chosen[THEME_COLUMN] = [{"code": theme["name"], "label": theme["name"]} for theme in themes]
    chosen[SCALE_COLUMN] = [{"code": value, "label": short} for value, short, _ in EXPERT_SCALE]
    return chosen


def fingerprint(text: Any) -> str:
    """Отпечаток текста строки: тот, что в манифесте, должен совпасть с текстом в таблице эксперта."""
    return hashlib.sha256(("" if text is None else str(text)).encode("utf-8")).hexdigest()


def open_book(path: Path) -> Any:
    import zipfile

    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile) as exc:
        raise PackError(f"{path.name}: не таблица xlsx ({exc})") from exc


def sheet_rows(book: Any, sheet: str, required: list[str], name: str) -> list[dict[str, Any]]:
    """Строки листа словарями по заголовку; строки без номера (ID) пропускаются. Нет листа или колонок — PackError."""
    if sheet not in book.sheetnames:
        raise PackError(f"{name}: нет листа «{sheet}»")
    rows = list(book[sheet].iter_rows(values_only=True))
    header = [text(cell) for cell in rows[0]] if rows else []
    absent = [column for column in required if column not in header]
    if absent:
        raise PackError(f"{name}: нет колонок {', '.join(absent)}")
    parsed = []
    for values in rows[1:]:
        cells = dict(zip(header, values))
        if text(cells.get("ID")):
            parsed.append({**cells, "ID": text(cells.get("ID"))})
    return parsed


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()

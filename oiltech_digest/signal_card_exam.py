"""Экзамен агента «Карточка сигнала» (#98, спека #93): модель выходит в работу по порогам на разметке эксперта.

Набор — пакет эксперту (#94): манифест хранит замороженные входы — заголовок, текст, дату публикации,
издателя и его тип, — таблица хранит разметку эксперта. Агенту подаются входы манифеста, а не база: ручные
поправки с первоисточника есть только там. Строка разметки сопоставляется с манифестом по номеру строки
(ID таблицы), а не по сигналу в базе: сигнал радара мог стать дублем.

Деление — 15 сигналов на настройку и 35 на экзамен по долям манифеста (наибольшие остатки, каждая доля в
обеих частях); оно пишется рядом с манифестом при первом прогоне и дальше не меняется. Инструкцию правят
только по строкам настройки: повторы экзамена на тех же строках превращали бы их в настройку.

Каждый сигнал — три прогона. Ответ проходит те же проверки кодом, что на ядре (signal_card.card_values),
— мерится то, что дойдёт до карточки. Модели — лестницей по цене: первая, прошедшая пороги
(signal_card_exam_score), — решение. Отчёт и решение пишутся рядом с манифестом, вне репозитория: он публичный.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
from typing import Any

from oiltech_digest import config, signal_card, signal_card_exam_baseline, signal_card_exam_report
from oiltech_digest.signal_card import _plain
from oiltech_digest.signal_card_exam_score import CompanyNames, baseline_metrics, evaluate, judged
from oiltech_digest.signal_card_pack import (LISTS_SHEET, MANIFEST_NAME, MARKUP_COLUMNS, MARKUP_SHEET, REPO_ROOT,
                                             UNDEFINED, _rank, load_dictionaries, publication_text)

# Лестница по цене (спека #93): сначала nano, при провале — mini. Названные модели идут так же — по цене.
MODELS = ("gpt-5.4-nano", "gpt-5.4-mini")
# Прогонов на сигнал экзамена (спека); на настройке — сколько угодно.
RUNS = 3
TUNING_SHARE = 15 / 50
SPLIT_NAME = "signal_card_exam_split.json"
# Допуски — ячейки, где текст и правила законно допускают второй ответ: {ID: {атрибут: [ответы]}}; у
# компаний — необязательные участники. Утверждает владелец; ключи на «_» и «note» — пояснения.
TOLERANCES_NAME = "signal_card_exam_tolerances.json"
# Алиасы экзамена — для названий, которые эксперт записал не так, как в тексте: [{name, aliases}], как у
# справочника компаний; к нему и добавляются (только для сравнения имён — агенту не передаются).
ALIASES_NAME = "signal_card_exam_aliases.json"
# Значения судьи радара у радарных строк набора — выгрузка с ядра (signal-card-exam-judge): {ID: {event_date,
# oilfield_relevance}}. Экзамен идёт там, где доступна модель, а база — на ядре.
JUDGE_NAME = "signal_card_exam_judge.json"
# Колонка таблицы разметки → атрибут; справочник значения — у выбора из списка.
LABEL_COLUMNS = {"Дата события": "event_date", "Источник даты события": "event_date_source",
                 "Компании-участники": "companies", "Тип события": "event_type", "Регион": "region",
                 "Класс технологии": "technology_class", "Зрелость": "maturity",
                 "Связь с нефтесервисом": "oilfield_relevance"}
LABEL_DICTIONARIES = {LABEL_COLUMNS[title]: key for title, key in MARKUP_COLUMNS if key}


class ExamError(ValueError):
    pass


def run(markup_path: str | Path, *, manifest_path: str | Path | None = None, models: tuple[str, ...] = MODELS,
        runs: int = RUNS, part: str = "exam") -> dict[str, Any]:
    """Экзамен по лестнице моделей: отчёт и решение — рядом с манифестом."""
    _check_runs(runs, part)
    exam = load(markup_path, manifest_path)
    split = _split(exam, write=True)
    items = _part_items(exam, split, part)
    dictionaries = load_dictionaries()
    names = CompanyNames(dictionaries["companies"], exam["aliases"])
    tried: list[dict[str, Any]] = []
    for model in _by_price(models):
        tried.append(_model_scores(model, _answers(model, items, runs, dictionaries), items, dictionaries, names))
        print(f"{model}: " + ("пороги пройдены" if tried[-1]["passed"]
                              else "не прошла: " + ", ".join(tried[-1]["lenient"]["failed"])))
        if _decision(tried, part):
            break
    settings = {"part": part, "runs": runs, "split": split, "reasoning": config.SIGNAL_CARD_REASONING,
                "prompt_version": signal_card.PROMPT_VERSION}
    return _report(exam, items, dictionaries, names, tried, settings)


def rescore(report_path: str | Path, markup_path: str | Path, *,
            manifest_path: str | Path | None = None) -> dict[str, Any]:
    """Пересчёт сохранённого отчёта без вызовов модели: те же ответы — по нынешним проверкам ядра
    (card_values), допускам, алиасам и выгрузке судьи; новый отчёт — рядом, с пометкой, откуда ответы. Ответы
    относятся к текстам своего манифеста и к своей версии справочников: с другими пересчёт отказывает."""
    source = Path(report_path).expanduser().resolve()
    try:
        saved = signal_card_exam_report.read(source)
    except ValueError as exc:
        raise ExamError(str(exc)) from exc
    exam = load(markup_path, manifest_path)
    dictionaries = load_dictionaries()
    if saved["manifest"]["sha256"] != exam["manifest"]["sha256"]:
        raise ExamError(f"{source.name}: ответы получены на другом манифесте — тексты другие, пересчёт неверен")
    if saved["dictionaries_version"] != dictionaries.get("version"):
        raise ExamError(f"{source.name}: ответы на справочниках {saved['dictionaries_version']}, сейчас "
                        f"{dictionaries.get('version')} — пересчёт неверен, нужен новый прогон")
    items = _part_items(exam, saved["split"], saved["part"])
    names = CompanyNames(dictionaries["companies"], exam["aliases"])
    # Цена и токены — исходного прогона: новых вызовов нет, а ставки с тех пор могли поменяться.
    tried = [{**_model_scores(one["model"], one["answers"], items, dictionaries, names),
              **{key: one[key] for key in ("calls", "input_tokens", "output_tokens", "cost_usd")}}
             for one in saved["models"]]
    settings = {key: saved[key] for key in ("part", "runs", "split", "reasoning", "prompt_version")}
    decided = saved.get("decision")
    return _report(exam, items, dictionaries, names, tried, settings, rescored_from=source.name,
                   markup_changed=saved["markup"]["sha256"] != exam["markup"]["sha256"],
                   original_decision=decided.get("model") if isinstance(decided, dict) else None)


def _part_items(exam: dict[str, Any], split: dict[str, Any], part: str) -> list[dict[str, Any]]:
    ids = set(split[part])
    return [item for item in exam["items"] if item["id"] in ids]


def _decision(tried: list[dict[str, Any]], part: str) -> str | None:
    """Решение экзамена — первая по цене модель, прошедшая пороги; настройка модель не выбирает."""
    return next((one["model"] for one in tried if one["passed"]), None) if part == "exam" else None


def _model_scores(model: str, answers: list[dict[str, Any]], items: list[dict[str, Any]],
                  dictionaries: dict[str, Any], names: CompanyNames) -> dict[str, Any]:
    """Ответы модели против разметки — строго и с допусками; цена — по токенам ответов."""
    finals = _finals(answers, items, dictionaries)
    scored = {view: evaluate(items, finals, names, lenient=view == "lenient") for view in ("strict", "lenient")}
    return {"model": model, **_spend(answers), "answers": answers, **scored,
            "without_card": sum(final is None for one in finals.values() for final in one),
            "passed": scored["lenient"]["passed"]}


def _report(exam: dict[str, Any], items: list[dict[str, Any]], dictionaries: dict[str, Any], names: CompanyNames,
            tried: list[dict[str, Any]], settings: dict[str, Any], **extra: Any) -> dict[str, Any]:
    """Отчёт прогона или пересчёта (settings — часть, прогоны, деление, рассуждение и версия инструкции) — рядом
    с манифестом, JSON и читаемый .md (signal_card_exam_report.write)."""
    part, decision = settings["part"], _decision(tried, settings["part"])
    report = {
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "part": part,
        **{key: settings[key] for key in ("runs", "reasoning", "prompt_version")},
        "dictionaries_version": dictionaries.get("version"), "manifest": exam["manifest"], "markup": exam["markup"],
        "files": exam["files"], "split": {key: settings["split"][key] for key in ("tuning", "exam")},
        "models": tried,
        "baseline": {view: baseline_metrics(items, signal_card_exam_baseline.finals(items, dictionaries), names,
                                            lenient=view == "lenient") for view in ("strict", "lenient")},
        "judge": judged(exam["items"], exam["judge"]),
        "decision": {"model": decision, "reason": signal_card_exam_report.reason(decision, tried) if part == "exam"
                     else "настройка: модель выбирает только экзамен"},
        **extra,
    }
    path, readable = signal_card_exam_report.write(exam["folder"], report)
    return {**report, "report": str(path), "readable": str(readable)}


def plan(markup_path: str | Path, *, manifest_path: str | Path | None = None, models: tuple[str, ...] = MODELS,
         runs: int = RUNS, part: str = "exam") -> dict[str, Any]:
    """Сухой прогон: число вызовов и цена по моделям; ничего не пишет и не вызывает модель. Цена — по длине
    запросов набора (signal_card.estimate)."""
    _check_runs(runs, part)
    exam = load(markup_path, manifest_path)
    split = _split(exam, write=False)
    items = _part_items(exam, split, part)
    dictionaries = load_dictionaries()
    chars = sum(signal_card.request_chars(_model_item(index, item), dictionaries) for index, item in enumerate(items))
    priced = [signal_card.estimate(chars * runs, len(items) * runs, model) for model in _by_price(models)]
    stratum_of = {item["id"]: item["stratum"] for item in exam["items"]}
    return {"part": part, "signals": len(items), "runs": runs, "rows": len(exam["items"]), "split_new": split["new"],
            "split": {key: dict(sorted(Counter(stratum_of[one] for one in split[key]).items()))
                      for key in ("tuning", "exam")},
            "models": priced}


def _check_runs(runs: int, part: str) -> None:
    """Экзамен — RUNS прогонов на сигнал (спека): меньше — устойчивость не измерить, и решение о модели было бы
    принято не так, как задано. Настройка — от одного."""
    if runs < 1 or (part == "exam" and runs != RUNS):
        raise ExamError(f"--runs {runs}: на экзамене — {RUNS} прогона на сигнал, на настройке — от одного")


def _by_price(models: tuple[str, ...]) -> list[str]:
    """Лестница — от дешёвой модели к дорогой, в каком порядке модели ни назвали: решение — самая дешёвая из прошедших."""
    return sorted(models, key=config.price_for_model)


def judge_values(manifest: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """Выгрузка с ядра (там база): значения судьи у радарных строк манифеста — по номеру строки."""
    from oiltech_digest.db import signal_card_store

    numbers = {item["id"]: int(item["signal_id"]) for item in manifest.get("items") or []
               if item.get("origin") == "radar" and item.get("signal_id")}
    found = signal_card_store.judge_values(list(numbers.values()))
    return {number: {"signal_id": signal_id, **found[signal_id]}
            for number, signal_id in numbers.items() if signal_id in found}


# --- Набор ---------------------------------------------------------------------------------------


def load(markup_path: str | Path, manifest_path: str | Path | None = None) -> dict[str, Any]:
    """Манифест и разметка эксперта: строки набора с входами агента и метками эксперта."""
    markup = Path(markup_path).expanduser().resolve()
    manifest_file = Path(manifest_path).expanduser().resolve() if manifest_path else markup.with_name(MANIFEST_NAME)
    folder = manifest_file.parent
    if folder == REPO_ROOT or REPO_ROOT in folder.parents:
        raise ExamError(f"{folder} — внутри репозитория, а он публичный: пакет эксперту и отчёт — вне его")
    manifest = json.loads(manifest_file.read_text(encoding="utf-8"))
    if not isinstance(manifest, dict) or not isinstance(manifest.get("items"), list) or not all(
            isinstance(item, dict) and {"id", "stratum", "text_sha256"} <= item.keys() for item in manifest["items"]):
        raise ExamError(f"{manifest_file.name}: не манифест пакета эксперту — нужен список items с id, stratum и "
                        "text_sha256")
    dictionaries = load_dictionaries()
    labels = _read_markup(markup, dictionaries)
    listed = [item["id"] for item in manifest["items"]]
    missing, extra = sorted(set(listed) - set(labels)), sorted(set(labels) - set(listed))
    if missing or extra:
        raise ExamError("разметка не от этого манифеста: " + "; ".join(
            part for part in (f"нет строк {', '.join(missing)}" if missing else "",
                              f"лишние строки {', '.join(extra)}" if extra else "") if part))
    tolerated = _tolerances(folder / TOLERANCES_NAME, set(listed), dictionaries)
    items = []
    for item in manifest["items"]:
        row = labels[item["id"]]
        if row["text_sha256"] != item["text_sha256"]:
            raise ExamError(f"{item['id']}: текст в таблице не тот, что в манифесте — таблица от другого пакета")
        items.append({**_inputs(item), "labels": row["labels"], "tolerated": tolerated.get(item["id"], {})})
    return {"folder": folder, "seed": manifest.get("seed"), "items": items, "aliases": _aliases(folder / ALIASES_NAME),
            "judge": _judge_file(folder / JUDGE_NAME),
            "manifest": _file_mark(manifest_file), "markup": _file_mark(markup),
            "files": {name: _mark_if_exists(folder / file) for name, file in (
                ("tolerances", TOLERANCES_NAME), ("aliases", ALIASES_NAME), ("judge", JUDGE_NAME))}}


def _aliases(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    entries = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(entries, list) or not all(
            isinstance(entry, dict) and isinstance(entry.get("name"), str) and isinstance(entry.get("aliases"), list)
            and all(isinstance(alias, str) for alias in entry["aliases"]) for entry in entries):
        raise ExamError(f"{path.name}: нужен список {{\"name\": …, \"aliases\": [...]}}, как у справочника компаний")
    return entries


def _judge_file(path: Path) -> dict[str, dict[str, Any]] | None:
    """Выгрузка судьи с ядра (signal-card-exam-judge): {ID: {event_date, oilfield_relevance}} → значения в форме
    карточки. Проверяется при загрузке набора, до вызовов модели: битый файл не должен сорвать оплаченный прогон."""
    if not path.exists():
        return None
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not all(isinstance(judged, dict) for judged in data.values()):
        raise ExamError(f"{path.name}: нужен объект {{ID строки: {{event_date, oilfield_relevance}}}} — выгрузка "
                        "signal-card-exam-judge")
    values = {}
    for number, judged in data.items():
        day, relevance = judged.get("event_date"), judged.get("oilfield_relevance")
        if relevance is not None and not isinstance(relevance, str):
            raise ExamError(f"{path.name}, {number}: связь с нефтесервисом {relevance!r} — не строка")
        try:
            values[number] = {"event_date": date.fromisoformat(day) if day else None, "oilfield_relevance": relevance}
        except (TypeError, ValueError) as exc:
            raise ExamError(f"{path.name}, {number}: дата события {day!r} — не ГГГГ-ММ-ДД") from exc
    return values


def _mark_if_exists(path: Path) -> dict[str, str] | None:
    return _file_mark(path) if path.exists() else None


def _tolerances(path: Path, ids: set[str], dictionaries: dict[str, Any]) -> dict[str, dict[str, list[Any]]]:
    """Список допусков: значения — в той же форме, что метки эксперта (справочники, даты ГГГГ-ММ-ДД)."""
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ExamError(f"{path.name}: нужен объект {{ID строки: {{атрибут: [ответы]}}}}")
    problems: list[str] = []
    tolerated: dict[str, dict[str, list[Any]]] = {}
    for number, cells in data.items():
        if number.startswith("_"):
            continue
        if number not in ids:
            problems.append(f"{number} — нет в манифесте")
            continue
        if not isinstance(cells, dict):
            problems.append(f"{number} — нужен объект {{атрибут: [ответы]}}")
            continue
        row: dict[str, list[Any]] = {}
        for key, values in cells.items():
            if key == "note":
                continue
            if key not in LABEL_COLUMNS.values():
                problems.append(f"{number}: атрибута {key!r} нет в разметке")
                continue
            values = values if isinstance(values, list) else [values]
            if key == "companies":
                row[key] = [" ".join(str(value).split()) for value in values if str(value).strip()]
            else:
                row[key] = [_label(key, value, dictionaries, problems, number) for value in values]
        tolerated[number] = row
    if problems:
        raise ExamError(f"{path.name} с ошибками: " + "; ".join(problems))
    return tolerated


def _file_mark(path: Path) -> dict[str, str]:
    return {"file": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def _inputs(item: dict[str, Any]) -> dict[str, Any]:
    """Входы агента из манифеста — так же, как ядро готовит главную публикацию (signal_card._publication)."""
    published = item.get("published")
    return {
        "id": item["id"], "stratum": item["stratum"], "origin": item.get("origin"),
        "title": publication_text(item.get("title"))[:500], "text": publication_text(item.get("text")),
        "published": date.fromisoformat(published) if published else None,
        "publisher": item.get("publisher") or "", "publisher_type": item.get("publisher_type") or "",
        "theme": item.get("theme") or "", "category": item.get("category") or UNDEFINED,
    }


def _read_markup(path: Path, dictionaries: dict[str, Any]) -> dict[str, dict[str, Any]]:
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if MARKUP_SHEET not in book.sheetnames:
        raise ExamError(f"{path.name}: нет листа «{MARKUP_SHEET}»")
    rows = list(book[MARKUP_SHEET].iter_rows(values_only=True))
    if not rows:
        raise ExamError(f"{path.name}: лист «{MARKUP_SHEET}» пуст")
    header = [str(cell or "").strip() for cell in rows[0]]
    absent = [name for name in ("ID", "Текст публикации", *LABEL_COLUMNS) if name not in header]
    if absent:
        raise ExamError(f"{path.name}: нет колонок {', '.join(absent)}")
    _check_lists(book, dictionaries)
    problems: list[str] = []
    labels: dict[str, dict[str, Any]] = {}
    for values in rows[1:]:
        cells = dict(zip(header, values))
        number = str(cells.get("ID") or "").strip()
        if not number:
            continue
        text = str(cells.get("Текст публикации") or "")
        labels[number] = {"text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(), "labels": {
            key: _label(key, cells.get(column), dictionaries, problems, number) for column, key in LABEL_COLUMNS.items()}}
    if problems:
        raise ExamError("разметка с ошибками: " + "; ".join(problems))
    return labels


def _check_lists(book: Any, dictionaries: dict[str, Any]) -> None:
    """Агент выбирает из тех же списков, из которых выбирал эксперт (скрытый лист выпадающих списков):
    другие значения справочника сделали бы экзамен нечестным, даже если метки остались допустимыми."""
    if LISTS_SHEET not in book.sheetnames:
        raise ExamError(f"в таблице нет листа «{LISTS_SHEET}»: не с чем сверить справочники")
    columns = list(zip(*book[LISTS_SHEET].iter_rows(values_only=True)))
    shown = {column[0]: [value for value in column[1:] if value not in (None, UNDEFINED)] for column in columns}
    changed = [key for key in sorted(set(LABEL_DICTIONARIES.values()))
               if shown.get(key) != list(dictionaries[key]["values"])]
    if changed:
        raise ExamError("справочники агента не те, по которым размечал эксперт: " + ", ".join(changed)
                        + " — экзамен на них ничего не скажет")


def _label(key: str, raw: Any, dictionaries: dict[str, Any], problems: list[str], number: str) -> Any:
    text = "" if raw is None else str(raw).strip()
    if key == "event_date":
        if isinstance(raw, datetime):
            return raw.date()
        if isinstance(raw, date):
            return raw
        if _plain(text) in ("", UNDEFINED):
            return None
        try:
            return date.fromisoformat(text[:10])
        except ValueError:
            problems.append(f"{number}: дата события {text!r} — не ГГГГ-ММ-ДД")
            return None
    if key == "companies":
        names = [" ".join(name.split()) for name in text.split(";")]
        return [name for name in names if name and _plain(name) != UNDEFINED]
    allowed = [*dictionaries[LABEL_DICTIONARIES[key]]["values"], UNDEFINED]
    found = next((item for item in allowed if _plain(item) == _plain(text)), None)
    if found is None:
        problems.append(f"{number}: {key} {text!r} — нет в справочнике")
    return found



# --- Деление -------------------------------------------------------------------------------------


def _split(exam: dict[str, Any], *, write: bool) -> dict[str, Any]:
    """Деление на настройку и экзамен: из файла рядом с манифестом, а при первом прогоне — новое."""
    path = exam["folder"] / SPLIT_NAME
    ids = sorted(item["id"] for item in exam["items"])
    if path.exists():
        saved = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(saved, dict) or not all(isinstance(saved.get(key), list) for key in ("tuning", "exam")):
            raise ExamError(f"{path.name}: испорчен — нужны списки tuning и exam")
        if sorted(saved["tuning"] + saved["exam"]) != ids:
            raise ExamError(f"{path.name} — деление другого набора: манифест сменился после первого прогона")
        return {**saved, "new": False}
    by_stratum: dict[str, list[str]] = defaultdict(list)
    for item in exam["items"]:
        by_stratum[item["stratum"]].append(item["id"])
    total = round(len(ids) * TUNING_SHARE)
    quotas = {stratum: len(group) * total / len(ids) for stratum, group in by_stratum.items()}
    counts = {stratum: int(quota) for stratum, quota in quotas.items()}
    ranked = sorted(by_stratum, key=lambda stratum: (counts[stratum] - quotas[stratum], -len(by_stratum[stratum]),
                                                     stratum))
    for stratum in ranked[:total - sum(counts.values())]:
        counts[stratum] += 1
    for stratum, group in by_stratum.items():
        if len(group) >= 2:  # каждая доля — в обеих частях
            counts[stratum] = min(max(counts[stratum], 1), len(group) - 1)
    tuning = {number for stratum, group in by_stratum.items()
              for number in sorted(group, key=lambda one: _rank("split", exam["seed"], one))[:counts[stratum]]}
    split = {"created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
             "rule": "наибольшие остатки по долям манифеста, порядок в доле — по хэшу зерна и номера",
             "tuning": [item["id"] for item in exam["items"] if item["id"] in tuning],
             "exam": [item["id"] for item in exam["items"] if item["id"] not in tuning]}
    if write:
        path.write_text(json.dumps(split, ensure_ascii=False, indent=2), encoding="utf-8")
    return {**split, "new": True}



# --- Прогоны -------------------------------------------------------------------------------------


def _model_item(index: int, item: dict[str, Any]) -> dict[str, Any]:
    keys = ("title", "text", "publisher", "publisher_type", "theme", "category")
    return {"signal_id": index, "publication": item["id"], **{key: item[key] for key in keys},
            "published": item["published"].isoformat() if item["published"] else None}


def _answers(model: str, items: list[dict[str, Any]], runs: int, dictionaries: dict[str, Any]) -> list[dict[str, Any]]:
    """Ответы модели: строка на (прогон, сигнал). Сбой вызова — одна повторная попытка, как у шага
    планировщика; обе попытки оплачены."""
    payload: dict[str, Any] = {"model": model, "reasoning": config.SIGNAL_CARD_REASONING, "dictionaries": dictionaries,
                               "signals": [_model_item(index, item) for index, item in enumerate(items)]}
    answers: list[dict[str, Any]] = []
    for number in range(1, runs + 1):
        cards = signal_card.process_payload(payload)["cards"]
        paid = [_tokens(card) for card in cards]
        failed = [index for index, card in enumerate(cards) if _failed(card, items[index], dictionaries)]
        if failed:
            retry = signal_card.process_payload({**payload, "signals": [payload["signals"][i] for i in failed]})
            for index, card in zip(failed, retry["cards"]):
                cards[index] = {**card, "attempts": 2}
                tokens_in, tokens_out = _tokens(card)
                paid[index] = (paid[index][0] + tokens_in, paid[index][1] + tokens_out)
        for index, card in enumerate(cards):
            answers.append({"run": number, "id": items[index]["id"], "answer": card.get("answer"),
                            "error": card.get("error"), "model": card.get("model"),
                            "attempts": card.get("attempts", 1), "input_tokens": paid[index][0],
                            "output_tokens": paid[index][1]})
        errors = sum(_failed(card, items[index], dictionaries) for index, card in enumerate(cards))
        print(f"{model}: прогон {number}/{runs} — вызовов {len(cards) + len(failed)}, без карточки {errors}")
    return answers


def _tokens(card: dict[str, Any]) -> tuple[int, int]:
    return int(card.get("input_tokens") or 0), int(card.get("output_tokens") or 0)


def _failed(card: dict[str, Any], item: dict[str, Any], dictionaries: dict[str, Any]) -> bool:
    return signal_card.card_values(card, item, dictionaries)[1] is not None


def _finals(answers: list[dict[str, Any]], items: list[dict[str, Any]],
            dictionaries: dict[str, Any]) -> dict[str, list[dict[str, Any] | None]]:
    """Значения карточки по прогонам — как их записало бы ядро; без карточки — None."""
    by_id = {item["id"]: item for item in items}
    finals: dict[str, list[dict[str, Any] | None]] = {item["id"]: [] for item in items}
    for answer in answers:
        values, _ = signal_card.card_values(answer, by_id[answer["id"]], dictionaries)
        finals[answer["id"]].append(values)
    return finals


def _spend(answers: list[dict[str, Any]]) -> dict[str, Any]:
    tokens_in = sum(answer["input_tokens"] for answer in answers)
    tokens_out = sum(answer["output_tokens"] for answer in answers)
    cost = sum((answer["input_tokens"] * price[0] + answer["output_tokens"] * price[1]) / 1_000_000
               for answer in answers for price in [config.price_for_model(answer["model"])])
    return {"calls": sum(answer["attempts"] for answer in answers), "input_tokens": tokens_in,
            "output_tokens": tokens_out, "cost_usd": round(cost, 4)}

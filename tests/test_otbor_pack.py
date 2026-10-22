"""Пакет эксперту для «Отбора» (#129, шов 2 спеки #128): справочники, черновики линеек, разметка вслепую.

Шов — команда `otbor-expert-pack` на изолированной базе; синтетика — otbor_pack_support. Настоящий пакет
уходит эксперту вне репозитория (он публичный)."""

from __future__ import annotations

from collections import Counter, defaultdict
import hashlib
from pathlib import Path
import re

import openpyxl
import pytest

from oiltech_digest import otbor_pack, otbor_pack_format as pack_format
from oiltech_digest.db import repository
from oiltech_digest.scoring_profiles import SCORING_PRESETS
from tests.otbor_pack_support import (MACHINE, THEMES, UNTIL, article, build_corpus, manifest, markup, material,
                                      pack_args, run_pack)

DASHES = re.compile("[–—]")
REASONS = ["Нефтегаз без связи с нефтесервисом", "Другая отрасль без переноса", "Происшествия, война, политика",
           "Корпоративное и реклама", "Старое событие", "Разные события в одной находке"]


@pytest.fixture()
def corpus(isolated_db, tmp_path):
    return build_corpus(tmp_path)


def _strings(value) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        return [text for one in value.values() for text in _strings(one)]
    if isinstance(value, list):
        return [text for one in value for text in _strings(one)]
    return []


def _rows(out: Path) -> dict[str, dict[str, object]]:
    _, _, header, rows = markup(out)
    return {row[0].value: dict(zip(header, (cell.value for cell in row))) for row in rows}


def _sheet_text(book, name: str) -> str:
    return " ".join(str(value) for row in book[name].iter_rows(values_only=True) for value in row if value is not None)


def test_dictionaries_and_rubric_drafts_are_versioned_data():
    dictionaries = pack_format.load_dictionaries()

    assert dictionaries["version"]
    labels = {key: [value["label"] for value in dictionaries[key]["values"]]
              for key in ("industry", "oilfield_relevance", "decision", "reject_reason", "category")}
    assert labels["industry"] == ["нефтесервис", "нефтегаз", "другая отрасль"]
    assert labels["oilfield_relevance"] == ["прямая", "переносимая", "нет связи"]
    assert labels["category"] == ["бизнес", "технология", "другое"]
    assert labels["reject_reason"] == REASONS
    for key in labels:
        codes = [value["code"] for value in dictionaries[key]["values"]]
        assert len(set(codes)) == len(codes) and len(set(labels[key])) == len(labels[key]), key
        assert dictionaries[key]["title"] and dictionaries[key]["rule"], key
    for key in labels:
        assert all(value["rule"] and value["examples"] for value in dictionaries[key]["values"]), key
    # Граница прямой связи и правило переносимой — то, что эксперт утверждает.
    direct, transferable, _ = dictionaries["oilfield_relevance"]["values"]
    assert "программы бурения" in direct["rule"] and "Котировки" in direct["rule"]
    assert "горнодобыча" in transferable["rule"]
    # Причина отказа главнее связи: старое событие и разные события отклоняются и при прямой связи (спека #128).
    assert "главнее" in dictionaries["decision"]["rule"]
    # Старое событие у находки без даты публикации — от даты поступления.
    old_event = next(value for value in dictionaries["reject_reason"]["values"] if value["code"] == "old_event")
    assert "даты поступления" in old_event["rule"]

    rubrics = pack_format.load_rubrics()
    assert rubrics["version"]
    assert [(band["code"], band["min"], band["max"]) for band in rubrics["bands"]] == [
        ("strong", 80, 100), ("useful", 65, 79), ("watch", 40, 64), ("weak", 0, 39)]
    # Линейка у каждого из десяти критериев двух профилей заказчика.
    for profile, criteria in SCORING_PRESETS["viktor"].items():
        assert [one["name"] for one in rubrics["profiles"][profile]] == [one["name"] for one in criteria]
        for criterion in rubrics["profiles"][profile]:
            assert set(criterion["bands"]) == {"strong", "useful", "watch", "weak"}, criterion["name"]
            assert all(band["rule"] and band["example"] for band in criterion["bands"].values()), criterion["name"]
    # Тексты уходят эксперту: без длинного и среднего тире.
    for text in _strings(dictionaries) + _strings(rubrics):
        assert not DASHES.search(text), text


def test_dry_run_prints_the_pack_by_groups_and_writes_nothing(corpus, tmp_path, capsys):
    out = tmp_path / "pack"

    run_pack(*pack_args(corpus, out))

    printed = capsys.readouterr().out
    assert "строк 116" in printed
    for part in ("прежняя разметка карточки 6", "отклонённые гейтом статьи ленты 50",
                 "принятые гейтом статьи ленты 30", "отсеянные и скрытые находки радара 18",
                 "находки с несколькими ссылками 10", "материалы коллег 2",
                 "брак судьи 6", "брак ревью пачки 5", "смешанные ссылки 4", "старые события 3",
                 "с 2026-09-22 по 2026-10-22"):
        assert part in printed, part
    assert "Не хватает" not in printed and "Мешает записи" not in printed
    assert not out.exists()


def test_apply_writes_every_group_blind_and_freezes_the_exam_inputs(corpus, tmp_path):
    out = tmp_path / "pack"

    run_pack(*pack_args(corpus, out, "--apply"))

    data = manifest(out)
    items = data["items"]
    groups = defaultdict(set)
    for item in items:
        groups[item["group"]].add(item["id"])
    assert groups["card_markup"] == {item["id"] for item in corpus["card"]["items"]}
    assert groups["feed_rejected"] == {f"article:{one}" for one in corpus["rejected"]}
    assert groups["feed_accepted"] == {f"article:{one}" for one in corpus["accepted"]}
    assert groups["radar_filtered"] == {f"signal:{one}" for ids in corpus["filtered"].values() for one in ids}
    assert groups["radar_links"] == {f"signal:{one}" for one in corpus["links"]}
    assert groups["materials"] == {f"material:{one}" for one in corpus["materials"]}
    assert Counter(item["kind"] for item in items if item["group"] == "radar_filtered") == {
        "judge": 6, "review": 5, "mixed": 4, "old": 3}
    listed = {item["id"] for item in items}
    excluded = ({f"article:{one}" for one in corpus["excluded_feed"]}
                | {f"signal:{one}" for one in corpus["excluded_radar"]}
                | {f"material:{one}" for one in corpus["excluded_materials"]})
    assert not listed & excluded
    assert len(listed) == len(items) == 116
    assert data["window"] == {"since": "2026-09-22", "until": UNTIL, "days": 30}
    # Версии справочников и линеек — в рецепте пакета: экзамен и агент сверяют их со своими.
    assert data["recipe"]["dictionaries_version"] == pack_format.load_dictionaries()["version"]
    assert data["recipe"]["rubrics_version"] == pack_format.load_rubrics()["version"]
    # Манифест хранит входы экзамена: заголовок, текст, дату публикации, издателя, ссылки находки.
    assert all({"title", "text", "published", "publisher", "links", "text_sha256"} <= set(item) for item in items)
    link = next(item for item in items if item["id"] == f"material:{min(corpus['materials'])}")
    assert link["title"] == "Author title link" and link["text"].startswith("Author description of link")
    assert link["published"] == "2026-10-10" and link["publisher"] == "colleague.example.net"

    book, _, header, rows = markup(out)
    assert len(rows) == 116
    cells = " ".join(str(cell.value) for row in rows for cell in row if cell.value is not None)
    for value in MACHINE:
        assert value not in cells, value
    # Группа не видна ни колонкой, ни порядком строк: блоками строки выдавали бы решение гейта и радара.
    for label in otbor_pack.GROUPS.values():
        assert label not in cells and label not in header
    group_of = {item["id"]: item["group"] for item in items}
    order = [group_of[row[header.index("ID")].value] for row in rows]
    for group in ("feed_rejected", "feed_accepted", "radar_filtered", "radar_links"):
        positions = [index for index, value in enumerate(order) if value == group]
        assert positions != list(range(positions[0], positions[0] + len(positions))), group
    # Тот же текст, что получит агент: отпечаток манифеста — от текста в таблице.
    by_id = _rows(out)
    for item in items:
        text = by_id[item["id"]]["Текст публикации"]
        assert hashlib.sha256(text.encode("utf-8")).hexdigest() == item["text_sha256"]
    # Колонки разметки пусты; связь заполнена только у прежних строк — ответом самого эксперта.
    answer_columns = header[len(pack_format.GIVEN_COLUMNS):]
    for item_id, row in by_id.items():
        filled = {column for column in answer_columns if row[column] is not None}
        assert filled == ({"Связь с нефтесервисом"} if group_of[item_id] == "card_markup" else set()), item_id


def test_card_markup_rows_keep_what_the_expert_saw_and_his_relevance(corpus, tmp_path):
    out = tmp_path / "pack"

    run_pack(*pack_args(corpus, out, "--apply"))

    items = {item["id"]: item for item in manifest(out)["items"]}
    rows = _rows(out)
    for card in corpus["card"]["items"]:
        item = items[card["id"]]
        # Входы — из прежнего манифеста (там поправки с первоисточника), а не из базы.
        assert [item[key] for key in ("title", "text", "published", "publisher", "text_sha256")] == [
            card[key] for key in ("title", "text", "published", "publisher", "text_sha256")]
        assert rows[card["id"]]["Связь с нефтесервисом"] == corpus["card"]["relevance"][card["id"]]
    first = corpus["card"]["items"][0]["id"]
    assert items[first]["overrides"] == ["published"] and items[first]["override_note"]
    # Прежняя находка с несколькими ссылками получает их список — агент получит его так же.
    with_links, single = (items[f"signal:{one}"] for one in corpus["card_radar"])
    assert len(with_links["links"]) == 3 and single["links"] == []
    assert rows[with_links["id"]]["Ссылки находки"] and rows[single["id"]]["Ссылки находки"] is None


def test_links_of_a_find_are_its_input_without_the_search_model_values(corpus, tmp_path):
    out = tmp_path / "pack"

    run_pack(*pack_args(corpus, out, "--apply"))

    items = {item["id"]: item for item in manifest(out)["items"]}
    rows = _rows(out)
    many = items[f"signal:{corpus['links_many']}"]
    assert len(many["links"]) == otbor_pack.LINKS_LIMIT == 8
    # Ссылка поиска: издатель — сайт из адреса (имя записала модель), даты нет (её дала модель поиска).
    assert many["links"][0] == {"url": "https://www.radar0.example.com/links-many",
                                "title": "Page title links-many #0", "publisher": "radar0.example.com",
                                "published": None}
    # Ссылка на статью корпуса: заголовок, издатель и дата — статьи.
    assert many["links"][1] == {"url": "https://news.example.org/corpus-link", "title": "Publisher title corpus-link",
                                "publisher": "World Oil", "published": "2026-09-10"}
    cell = rows[many["id"]]["Ссылки находки"]
    assert cell.count("\n8. ") == 1 and "9. " not in cell and "Publisher title corpus-link" in cell
    # Главная публикация находки: издатель — сайт, даты публикации нет, ссылок одна — колонка пуста.
    judge = rows[f"signal:{corpus['filtered']['judge'][0]}"]
    assert judge["Издатель"] == "radar0.example.com" and judge["Дата публикации"] is None
    assert judge["Ссылки находки"] is None
    assert all(len(items[f"signal:{one}"]["links"]) == 2 for one in corpus["filtered"]["mixed"])


def test_received_date_anchors_the_age_of_an_event_without_a_publication_date(corpus, tmp_path):
    # У находки «режима ChatGPT» даты публикации нет (её дала модель поиска): возраст события эксперт и агент
    # считают от дня, когда публикация пришла на платформу. У прежних строк его нет в манифесте — берётся из базы.
    out = tmp_path / "pack"

    run_pack(*pack_args(corpus, out, "--apply"))

    items = {item["id"]: item for item in manifest(out)["items"]}
    rows = _rows(out)
    old = f"signal:{corpus['filtered']['old'][0]}"
    assert rows[old]["Дата публикации"] is None and rows[old]["Дата поступления"] == "2026-10-01"
    assert items[old]["received"] == "2026-10-01"
    feed = next(item for item in items.values() if item["group"] == "feed_rejected")
    assert feed["received"] == feed["published"] and rows[feed["id"]]["Дата поступления"] == feed["received"]
    card = {item["id"]: item for item in items.values() if item["group"] == "card_markup"}
    assert {card[f"article:{one}"]["received"] for one in corpus["card_feed"]} == {"2026-10-02"}
    assert {card[f"signal:{one}"]["received"] for one in corpus["card_radar"]} == {"2026-10-01"}
    assert all(card[f"article:{one}"]["language"] == "en" for one in corpus["card_feed"])
    assert items[f"material:{min(corpus['materials'])}"]["received"] == "2026-10-10"


def test_a_file_material_brings_the_text_of_its_file(corpus, tmp_path):
    # У файла во вход «Отбора» идёт текст файла (спека #128): эксперт размечает то же, что получит агент.
    import docx

    document = docx.Document()
    document.add_paragraph("Текст файла: протокол промысловых испытаний долота.")
    report = tmp_path / "report.docx"
    document.save(report)
    with repository.get_connection() as conn:
        readable, broken = material(conn, "file", kind="file"), material(conn, "broken", kind="file")
        conn.execute("UPDATE colleague_materials SET file_path = %s WHERE id = %s", (str(report), readable))
        conn.execute("UPDATE colleague_materials SET file_path = %s WHERE id = %s",
                     (str(tmp_path / "missing.pdf"), broken))
        conn.commit()
    out = tmp_path / "pack"

    run_pack(*pack_args(corpus, out, "--apply"))

    items = {item["id"]: item for item in manifest(out)["items"]}
    with_file, without = items[f"material:{readable}"], items[f"material:{broken}"]
    assert with_file["text"].startswith("Author description of file")
    assert "протокол промысловых испытаний долота" in with_file["text"] and with_file["file_note"] is None
    # Файл не прочитался: только описание автора, причина — в манифесте.
    assert without["text"].startswith("Author description of broken") and "не прочитан" in without["file_note"]


def test_columns_lists_rules_and_sheets_for_approval(corpus, tmp_path):
    out = tmp_path / "pack"

    run_pack(*pack_args(corpus, out, "--apply"))

    book, sheet, header, _ = markup(out)
    rubrics = pack_format.load_rubrics()
    scores = ([f"Бизнес: {one['name']}" for one in rubrics["profiles"]["business"]]
              + [f"Технология: {one['name']}" for one in rubrics["profiles"]["tech_radar"]])
    assert header == pack_format.GIVEN_COLUMNS + [
        "Отрасль", "Связь с нефтесервисом", "Брать или нет", "Причина отказа", "Категория", "Тема", *scores,
        "Почему такой балл", "Общая оценка", "Комментарий эксперта"]
    letter = {name: openpyxl.utils.get_column_letter(index + 1) for index, name in enumerate(header)}
    validated = {openpyxl.utils.get_column_letter(cell_range.min_col): validation.type
                 for validation in sheet.data_validations.dataValidation for cell_range in validation.sqref.ranges}
    for name in ("Отрасль", "Связь с нефтесервисом", "Брать или нет", "Причина отказа", "Категория", "Тема",
                 "Общая оценка"):
        assert validated.get(letter[name]) == "list", name
    for name in scores[:-1]:
        assert validated.get(letter[name]) == "whole", name
    assert letter["Технология: Свежесть"] not in validated  # число или «нет даты»
    for name in ("Почему такой балл", "Комментарий эксперта"):
        assert letter[name] not in validated, name
    lists = book[pack_format.LISTS_SHEET]
    assert lists.sheet_state == "hidden"
    shown = {value for row in lists.iter_rows(values_only=True) for value in row if value}
    assert set(THEMES) <= shown and "Выключенная тема" not in shown
    assert repository.SYSTEM_TAG_UNCLASSIFIED not in shown

    # Лист правил: у каждой колонки разметки правило и пример; значения справочников — с правилом и примерами; темы —
    # с описанием.
    rules = list(book[pack_format.RULES_SHEET].iter_rows(values_only=True))
    ruled = {row[0]: row[1] for row in rules[1:] if row[0] and row[1]}
    assert set(header[len(pack_format.GIVEN_COLUMNS):]) <= set(ruled) | set(scores)
    by_value = {row[2]: row for row in rules[1:] if row[2]}
    listed = {"Отрасль", "Связь с нефтесервисом", "Брать или нет", "Причина отказа", "Категория"}
    for column in set(ruled) - listed - {"Ссылки находки", "Дата поступления"}:
        assert "Пример" in ruled[column], column  # у колонок со справочником примеры — у значений
    for value in ("нефтесервис", "прямая", "брать", "не брать", *REASONS, "бизнес"):
        assert "Примеры:" in by_value[value][3], value
    assert by_value["Бурение"][3] == THEMES["Бурение"]
    assert "вынести в дайджест" in by_value["Сильный сигнал"][3]

    # На утверждение: шесть причин отказа и граница прямой связи; линейки — десять критериев.
    approval = list(book[pack_format.APPROVAL_SHEET].iter_rows(values_only=True))
    assert approval[0][:2] == ("Раздел", "Значение") and "Решение" in approval[0]
    decision = approval[0].index("Решение")
    approved = [row[1] for row in approval[1:]]
    assert approved == [*REASONS, "прямая", "переносимая"]
    assert all(row[decision] is None for row in approval[1:])
    lines = list(book[pack_format.RUBRICS_SHEET].iter_rows(values_only=True))
    assert lines[0][2:6] == ("80 и выше", "65-79", "40-64", "ниже 40")
    assert [row[1] for row in lines[1:]] == [name.split(": ", 1)[1] for name in scores]
    assert all(row[2] and row[3] and row[4] and row[5] and "Пример: " in row[2] for row in lines[1:])
    guide = _sheet_text(book, "Инструкция")
    assert pack_format.load_dictionaries()["version"] in guide and rubrics["version"] in guide
    # Тексты для эксперта — без длинного и среднего тире (кроме текстов издателей на листе разметки).
    for name in ("Инструкция", pack_format.RULES_SHEET, pack_format.APPROVAL_SHEET, pack_format.RUBRICS_SHEET):
        assert not DASHES.search(_sheet_text(book, name)), name


def test_missing_card_markup_or_a_short_group_is_reported_and_blocks_writing(corpus, tmp_path, capsys):
    out = tmp_path / "pack"

    run_pack("--out", str(out), "--until", UNTIL)

    printed = capsys.readouterr().out
    assert "Не хватает" in printed and "--card-manifest" in printed
    with pytest.raises(SystemExit) as refused:
        run_pack("--out", str(out), "--until", UNTIL, "--apply")
    assert "--card-manifest" in str(refused.value)
    # Окно, в котором нет ни статей ленты, ни находок: пакет на таком наборе экзамену ничего не даст.
    with pytest.raises(SystemExit) as refused:
        run_pack(*pack_args(corpus, out, "--until", "2026-09-01", "--apply"))
    assert "отклонённые гейтом статьи ленты 0 из 50" in str(refused.value)
    assert not out.exists()


def test_card_markup_from_another_pack_is_refused(corpus, tmp_path):
    book = openpyxl.load_workbook(corpus["card"]["markup"])
    book["Разметка"]["C2"] = "Текст другого пакета"
    other = tmp_path / "other.xlsx"
    book.save(other)
    out = tmp_path / "pack"

    with pytest.raises(SystemExit) as refused:
        run_pack(*pack_args(corpus, out, "--apply", "--card-markup", str(other)))

    assert "текст" in str(refused.value) and corpus["card"]["items"][0]["id"] in str(refused.value)
    assert not out.exists()


def test_rubrics_must_match_the_criteria_on_the_scoring_screen(corpus, tmp_path, capsys):
    with repository.get_connection() as conn:
        conn.execute("UPDATE scoring_criteria SET name = 'Новизна' WHERE profile = 'tech_radar'"
                     " AND name = 'Технологическая новизна'")
        conn.commit()
    out = tmp_path / "pack"

    run_pack(*pack_args(corpus, out))

    printed = capsys.readouterr().out
    assert "Мешает записи" in printed and "Технологическая новизна" in printed and "Новизна" in printed
    with pytest.raises(SystemExit):
        run_pack(*pack_args(corpus, out, "--apply"))
    assert not out.exists()


def test_same_seed_and_window_give_the_same_rows_and_never_overwrite(corpus, tmp_path):
    with repository.get_connection() as conn:
        source = conn.execute("SELECT id FROM sources WHERE name = 'World Oil'").fetchone()[0]
        for number in range(12):  # кандидатов больше квоты: выбор — по зерну
            article(conn, source, f"rejected-extra-{number}", "2026-10-07", relevant=False)
        conn.commit()
    first, second, other = tmp_path / "a", tmp_path / "b", tmp_path / "c"

    run_pack(*pack_args(corpus, first, "--apply", "--seed", "7"))
    with repository.get_connection() as conn:  # лента живёт дальше: новое за окном выборку не меняет
        article(conn, source, "rejected-later", "2026-10-25", relevant=False)
        conn.commit()
    run_pack(*pack_args(corpus, second, "--apply", "--seed", "7"))
    run_pack(*pack_args(corpus, other, "--apply", "--seed", "8"))

    def rejected(out):
        return {item["id"] for item in manifest(out)["items"] if item["group"] == "feed_rejected"}

    assert [item["id"] for item in manifest(first)["items"]] == [item["id"] for item in manifest(second)["items"]]
    assert len(rejected(first)) == 50 and rejected(first) != rejected(other)
    # Повтор в ту же папку не затирает таблицу, которую эксперт мог уже начать размечать.
    with pytest.raises(SystemExit):
        run_pack(*pack_args(corpus, first, "--apply", "--seed", "7"))


def test_pack_is_never_written_inside_the_repository(corpus):
    inside = Path(otbor_pack.__file__).resolve().parents[1] / "exports" / "otbor-pack"

    with pytest.raises(SystemExit):
        run_pack(*pack_args(corpus, inside, "--apply"))

    assert not inside.exists()
    with pytest.raises(SystemExit):
        run_pack("--until", "22.10.2026")

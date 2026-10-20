"""Пакет эксперту для карточки сигнала (#94): справочники на утверждение и таблица разметки.

Шов — команда `signal-card-expert-pack` на изолированной базе. Статьи и сигналы здесь
синтетические: настоящая таблица уходит эксперту вне репозитория (он публичный)."""

from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

import openpyxl
import pytest

from oiltech_digest import cli, signal_card_pack
from oiltech_digest.db import repository

THEMES = ["Бурение", "Добыча", "Геология", "Экология"]
# Значения, которые поставила машина: ни одно не должно попасть в таблицу эксперта.
MACHINE = [
    "CARD_SUMMARY", "CARD_TITLE_RU", "RELEVANCE_REASON", "TAG_RATIONALE",
    "JUDGE_TITLE", "JUDGE_SUMMARY", "JUDGE_WHY", "JUDGE_APPLICATION", "JudgeCo",
    "RESEARCH_SUMMARY", "JUDGE_VERDICT", "2026-07-03",
]
# Издатель может прислать что угодно: «=» в начале заголовка Excel счёл бы формулой, а
# управляющий символ openpyxl не пишет вовсе — выгрузка падала бы целиком.
HOSTILE_TITLE = '=HYPERLINK("https://evil.example") '
CONTROL = "\x0b"


def _utc(month, day):
    return datetime(2026, month, day, 12, tzinfo=timezone.utc)


def _source_and_tags(conn):
    source_id = conn.execute(
        "INSERT INTO sources (name, source_type, url, enabled, parse_strategy)"
        " VALUES ('World Oil', 'Media', 'https://news.example.org', TRUE, 'request') RETURNING id"
    ).fetchone()[0]
    roots = [conn.execute("INSERT INTO tags (name) VALUES (%s) RETURNING id", (name,)).fetchone()[0]
             for name in THEMES]
    return source_id, roots


def _seed_feed(conn, source_id, tag_ids, n, *, month=9, day=None, relevant=True, prefix="feed", text=None, score=75):
    ids = []
    for i in range(n):
        at = _utc(month, day or 1 + i % 27)
        article_id = conn.execute(
            "INSERT INTO articles (source_id, title, url, raw_text, published_at, collected_at)"
            " VALUES (%s, %s, %s, %s, %s, %s) RETURNING id",
            (source_id, f"{HOSTILE_TITLE}{prefix} article {i}", f"https://news.example.org/{prefix}/{i}",
             text or f"Body of {prefix} {i}.{CONTROL} " * 400, at, at),
        ).fetchone()[0]
        conn.execute(
            "INSERT INTO article_cards (article_id, summary, title_ru, relevant, relevance_reason)"
            " VALUES (%s, 'CARD_SUMMARY', 'CARD_TITLE_RU', %s, 'RELEVANCE_REASON')",
            (article_id, relevant),
        )
        conn.execute(
            "INSERT INTO article_tags (article_id, tag_id, confidence, rationale) VALUES (%s, %s, 0.9, 'TAG_RATIONALE')",
            (article_id, tag_ids[i % len(tag_ids)]),
        )
        if score is not None:
            conn.execute("INSERT INTO article_scores (article_id, total_score) VALUES (%s, %s)", (article_id, score))
        ids.append(article_id)
    return ids


def _mark(conn, article_ids, status):
    """Отметка аналитика в ленте: «в дайджест», «шум», «дубликат»."""
    user_id = conn.execute(
        "INSERT INTO users (email, password_salt, password_hash, role) VALUES (%s, 'x', 'x', 'user') RETURNING id",
        (f"{status}-{article_ids[0]}@example.test",),
    ).fetchone()[0]
    for article_id in article_ids:
        conn.execute("INSERT INTO article_marks (set_by_user_id, article_id, status) VALUES (%s, %s, %s)",
                     (user_id, article_id, status))


def _radar(key, category, theme, *, fetched=True, filtered=False, seen="2026-09-10", text=None):
    signal_id = repository.upsert_signal({
        "signal_key": key, "title": f"JUDGE_TITLE {key}", "title_ru": f"JUDGE_TITLE {key}", "theme": theme,
        "summary": "JUDGE_SUMMARY", "maturity": "watch", "score": 70, "companies": ["JudgeCo"],
    })
    repository.upsert_signal_evidence(signal_id, {
        "source_url": f"https://radar.example.com/{key}", "title": f"Page title {key}",
        "published_at": "2026-07-03", "extracted_fact": (text or f"Page text of {key}. " * 10) if fetched else "JUDGE_SUMMARY",
        "summary_ru": "RESEARCH_SUMMARY", "strength": 0.8,
        "raw_payload": {"evidence_source": "openai_web_search", "full_text_fetched": fetched,
                        "research_summary": "RESEARCH_SUMMARY"},
    })
    with repository.get_connection() as conn:
        conn.execute(
            "UPDATE signals SET signal_category = %s, event_date = '2026-07-03', why_interesting = 'JUDGE_WHY',"
            " oilfield_relevance = 'direct', oilfield_application = 'JUDGE_APPLICATION',"
            " verdict_reason = 'JUDGE_VERDICT', filter_stage = %s, first_seen_at = %s WHERE id = %s",
            (category, "judge" if filtered else None, seen, signal_id),
        )
        conn.commit()
    return signal_id


@pytest.fixture()
def corpus(isolated_db):
    with repository.get_connection() as conn:
        source_id, roots = _source_and_tags(conn)
        # Часть статей помечена дочерним тегом: тема — корневой.
        child = conn.execute("INSERT INTO tags (name, parent_id) VALUES ('Бурение: долота', %s) RETURNING id",
                             (roots[0],)).fetchone()[0]
        # Полезная статья — балл не ниже порога дайджеста или отметка «в дайджест»; шумная — балл ниже
        # порога или отметка «шум». Отметка человека сильнее балла.
        useful = _seed_feed(conn, source_id, roots + [child], 30)
        noise = _seed_feed(conn, source_id, roots, 10, score=30, prefix="noise")
        digest_low = _seed_feed(conn, source_id, roots, 1, score=20, prefix="digest-low")
        noise_high = _seed_feed(conn, source_id, roots, 1, score=95, prefix="noise-high")
        unscored = _seed_feed(conn, source_id, roots, 1, score=None, prefix="unscored")
        duplicate = _seed_feed(conn, source_id, roots, 1, score=80, prefix="duplicate")
        _mark(conn, digest_low, "digest")
        _mark(conn, noise_high, "noise")
        _mark(conn, duplicate, "duplicate")
        rejected = _seed_feed(conn, source_id, roots, 5, relevant=False, prefix="rejected")
        july = _seed_feed(conn, source_id, roots, 5, month=7, prefix="july")
        reprint = _seed_feed(conn, source_id, roots, 1, prefix="reprint")
        conn.execute("INSERT INTO article_reprints (article_id, primary_id) VALUES (%s, %s)",
                     (reprint[0], useful[0]))
        conn.commit()
    tech = [_radar(f"tech{i}", "technology", THEMES[i % 3]) for i in range(18)]
    business = [_radar(f"biz{i}", "business", THEMES[i % 2]) for i in range(8)]
    excluded = [
        _radar("filtered", "technology", THEMES[0], filtered=True),
        _radar("nofetch", "technology", THEMES[0], fetched=False),  # текст — сводка модели, не страница
    ]
    merged = _radar("merged", "business", THEMES[0])
    with repository.get_connection() as conn:
        conn.execute("UPDATE signals SET merged_into_signal_id = %s WHERE id = %s", (business[0], merged))
        conn.commit()
    return {
        "source": source_id, "tags": roots,
        "useful": set(useful + digest_low), "noise": set(noise + noise_high),
        "feed": set(useful + digest_low + noise + noise_high),
        "not_feed": set(rejected + july + reprint + unscored + duplicate),
        "tech": set(tech), "business": set(business), "not_radar": set(excluded + [merged]),
    }


def _run(*args):
    cli.main(["signal-card-expert-pack", *args])


def _manifest(out: Path):
    return json.loads((out / signal_card_pack.MANIFEST_NAME).read_text(encoding="utf-8"))


def _markup(out: Path):
    book = openpyxl.load_workbook(out / signal_card_pack.WORKBOOK_NAME)
    sheet = book[signal_card_pack.MARKUP_SHEET]
    rows = list(sheet.iter_rows())
    return book, sheet, [cell.value for cell in rows[0]], rows[1:]


def test_dictionaries_are_versioned_nonempty_and_unique():
    dictionaries = signal_card_pack.load_dictionaries()

    assert dictionaries["version"]
    for key in ("event_type", "region", "technology_class", "maturity", "date_source", "oilfield_relevance"):
        values = dictionaries[key]["values"]
        assert values and all(str(value).strip() for value in values), key
        assert len(values) == len(set(values)), key
        assert signal_card_pack.UNDEFINED not in values, key  # «не определено» добавляет таблица, не справочник
    companies = dictionaries["companies"]
    assert len(companies) >= 30
    names = [company["name"] for company in companies]
    assert len(names) == len(set(names))
    # Одно написание — одна компания: иначе алиас свёл бы двух участников к одному имени.
    spellings = [spelling.lower() for company in companies for spelling in [company["name"], *company["aliases"]]]
    assert len(spellings) == len(set(spellings))


def test_dry_run_prints_the_sample_and_writes_nothing(corpus, tmp_path, capsys):
    out = tmp_path / "pack"

    _run("--out", str(out), "--until", "2026-10-01")

    printed = capsys.readouterr().out
    assert "лента 30" in printed and "радар 20" in printed
    assert "бизнес" in printed and "технология" in printed
    for theme in THEMES:
        assert theme in printed
    assert "--until 2026-10-01" in printed  # чем повторить ту же выборку
    assert not out.exists()


def test_apply_writes_50_blind_signals_and_a_manifest(corpus, tmp_path):
    out = tmp_path / "pack"

    _run("--out", str(out), "--apply")

    items = _manifest(out)["items"]
    feed = {item["article_id"] for item in items if item["origin"] == "feed"}
    radar = {item["signal_id"] for item in items if item["origin"] == "radar"}
    assert len(items) == 50 and len(feed) == 30 and len(radar) == 20
    assert feed <= corpus["feed"] and not feed & corpus["not_feed"]
    # Эксперт размечает в основном полезное: 24 полезных и 6 шумных — намеренно, проверить «нет связи».
    strata = {stratum: {item["article_id"] for item in items if item.get("stratum") == stratum}
              for stratum in ("useful", "noise")}
    assert len(strata["useful"]) == 24 and strata["useful"] <= corpus["useful"]
    assert len(strata["noise"]) == 6 and strata["noise"] <= corpus["noise"]
    assert len(radar & corpus["tech"]) == 15 and len(radar & corpus["business"]) == 5
    # Строки вперемешку: блоками доли выдавали бы эксперту решение машины — шумные шли бы последними
    # статьями ленты, технологии — первыми сигналами радара.
    _, _, header, rows = _markup(out)
    stratum_of = {item["id"]: item["stratum"] for item in items}
    order = [stratum_of[row[header.index("ID")].value] for row in rows]
    for stratum in ("useful", "noise", "technology", "business"):
        positions = [index for index, value in enumerate(order) if value == stratum]
        assert positions != list(range(positions[0], positions[0] + len(positions))), stratum
    assert radar <= corpus["tech"] | corpus["business"] and not radar & corpus["not_radar"]
    assert radar & corpus["tech"] and radar & corpus["business"]
    assert {item["category"] for item in items} == {"бизнес", "технология"}
    assert {item["theme"] for item in items if item["origin"] == "feed"} == set(THEMES)
    assert _manifest(out)["dictionaries_version"] == signal_card_pack.load_dictionaries()["version"]

    book, sheet, header, rows = _markup(out)
    assert len(rows) == 50
    cells = " ".join(str(cell.value) for row in rows for cell in row if cell.value is not None)
    for value in MACHINE:
        assert value not in cells, value
    assert CONTROL not in cells
    assert all(cell.data_type != "f" for row in rows for cell in row)
    # Таблица и манифест — об одних и тех же входах: отпечаток — от текста в таблице.
    column = {name: index for index, name in enumerate(header)}
    by_id = {row[column["ID"]].value: row for row in rows}
    for item in items:
        text = by_id[item["id"]][column["Текст публикации"]].value
        assert hashlib.sha256(text.encode("utf-8")).hexdigest() == item["text_sha256"]
        if item["origin"] == "feed":
            assert len(text) == signal_card_pack.TEXT_LIMIT == 6000
            assert by_id[item["id"]][column["Заголовок"]].value.startswith("=HYPERLINK")

    # Разметка — вслепую: колонки атрибутов пусты; у справочных — выпадающий список.
    attributes = [name for name, _ in signal_card_pack.MARKUP_COLUMNS]
    assert all(row[column[name]].value is None for row in rows for name in attributes)
    validated = {}
    for validation in sheet.data_validations.dataValidation:
        for cell_range in validation.sqref.ranges:
            validated[openpyxl.utils.get_column_letter(cell_range.min_col)] = validation.type
    letter = {name: openpyxl.utils.get_column_letter(column[name] + 1) for name in header}
    for name in ("Тип события", "Регион", "Класс технологии", "Зрелость", "Источник даты события",
                 "Связь с нефтесервисом"):
        assert validated.get(letter[name]) == "list", name
    for name in ("Компании-участники", "Дата события"):
        assert letter[name] not in validated, name
    assert signal_card_pack.DICTIONARY_SHEET in book.sheetnames


def test_feed_signals_of_the_unified_store_stay_out_of_the_radar_part(corpus, tmp_path):
    # Единое хранилище (#95): принятая статья ленты — сигнал «бизнес» со статьёй-публикацией. Радарная
    # часть — только находки поиска, иначе сигналы ленты заняли бы квоту «бизнес».
    from oiltech_digest.db import feed_signals

    assert feed_signals.sync_feed_signals(collected_since="2026-01-01")["signals_created"] >= len(corpus["feed"])
    out = tmp_path / "pack"

    _run("--out", str(out), "--apply")

    radar = {item["signal_id"] for item in _manifest(out)["items"] if item["origin"] == "radar"}
    assert len(radar) == 20 and radar <= corpus["tech"] | corpus["business"]


def test_same_seed_and_window_give_the_same_sample_and_never_overwrite(corpus, tmp_path):
    first, second = tmp_path / "a", tmp_path / "b"
    window = ("--seed", "7", "--since", "2026-08-01", "--until", "2026-10-01")

    _run("--out", str(first), "--apply", *window)
    # Лента и радар живут дальше: новое за окном выборку не меняет.
    with repository.get_connection() as conn:
        _seed_feed(conn, corpus["source"], corpus["tags"], 12, month=10, day=5, prefix="later")
        conn.commit()
    _radar("later", "business", THEMES[3], seen="2026-10-05")
    _run("--out", str(second), "--apply", *window)

    def ids(out):
        return [item["id"] for item in _manifest(out)["items"]]

    assert ids(first) == ids(second)
    # Повтор в ту же папку не затирает таблицу, которую эксперт мог уже начать размечать.
    with pytest.raises(SystemExit):
        _run("--out", str(first), "--apply", *window)


def test_usefulness_comes_from_the_score_and_the_analysts_marks(isolated_db, capsys):
    # Владелец 15.10: в первой таблице половина статей ленты — не про нефтесервис (цены, газификация,
    # происшествия). Отбор пропускает всё «про ТЭК»; полезность — балл не ниже порога дайджеста или
    # отметка «в дайджест», а отметка «шум» делает статью шумной при любом балле.
    with repository.get_connection() as conn:
        source_id, roots = _source_and_tags(conn)
        _seed_feed(conn, source_id, roots, 1, score=75, prefix="strong")
        _mark(conn, _seed_feed(conn, source_id, roots, 1, score=20, prefix="digest-low"), "digest")
        _mark(conn, _seed_feed(conn, source_id, roots, 1, score=95, prefix="noise-high"), "noise")
        _seed_feed(conn, source_id, roots, 1, score=30, prefix="weak")
        _seed_feed(conn, source_id, roots, 1, score=None, prefix="unscored")  # балла ещё нет — не знаем
        _mark(conn, _seed_feed(conn, source_id, roots, 1, score=80, prefix="duplicate"), "duplicate")
        # Отметка одна на статью (#111): «в дайджест» при низком балле — полезная.
        _mark(conn, _seed_feed(conn, source_id, roots, 1, score=40, prefix="both"), "digest")
        # «Архив» ставят и «скрыть», и «убрать из дайджеста» — по нему не понять, полезна ли статья.
        _mark(conn, _seed_feed(conn, source_id, roots, 1, score=90, prefix="archived"), "archive")
        conn.commit()

    _run("--until", "2026-10-01")

    printed = capsys.readouterr().out
    assert "полезных статей ленты 3 из 24" in printed
    assert "шумных статей ленты 2 из 6" in printed


def test_a_short_stratum_is_not_filled_from_its_neighbour(isolated_db, tmp_path, capsys):
    # Добор из соседней доли тихо вернул бы прежний перекос: недостающий шум — полезными статьями,
    # недостающий бизнес — технологиями.
    with repository.get_connection() as conn:
        source_id, roots = _source_and_tags(conn)
        _seed_feed(conn, source_id, roots, 30)
        _seed_feed(conn, source_id, roots, 2, score=30, prefix="noise")
        conn.commit()
    for i in range(20):
        _radar(f"tech{i}", "technology", THEMES[i % 3])
    for i in range(2):
        _radar(f"biz{i}", "business", THEMES[0])
    out = tmp_path / "pack"

    _run("--until", "2026-10-01")

    lines = capsys.readouterr().out.splitlines()
    assert "Для экзамена не хватает: шумных статей ленты 2 из 6; бизнес-сигналов радара 2 из 5" in lines
    with pytest.raises(SystemExit):
        _run("--out", str(out), "--apply", "--until", "2026-10-01")
    assert not out.exists()


def test_block_pages_and_unreadable_texts_are_not_offered_to_the_expert(isolated_db, capsys):
    # С прода, первая выгрузка 15.10: докачка радара сохранила страницу блокировки вместо статьи,
    # а один сигнал радара — на китайском. Разметить их эксперт не может.
    blocked = "Доступ с вашего IP-адреса ограничен. Отключите VPN/прокси (если включены) и обновите страницу. " * 3
    chinese = "大庆油田采油八厂第四作业区407注采班中控室内，集输工端坐在操作台前，目光在数据大屏上快速扫过。" * 4
    with repository.get_connection() as conn:
        source_id, roots = _source_and_tags(conn)
        _seed_feed(conn, source_id, roots, 1, prefix="ok")
        _seed_feed(conn, source_id, roots, 1, prefix="blocked", text=blocked)
        _seed_feed(conn, source_id, roots, 1, prefix="chinese", text=chinese)
        conn.commit()
    _radar("ok", "technology", THEMES[0])
    _radar("blocked", "technology", THEMES[0], text=blocked)
    _radar("chinese", "business", THEMES[1], text=chinese)

    _run("--until", "2026-10-01")

    printed = capsys.readouterr().out
    assert "лента 1, радар 1" in printed


def test_incomplete_sample_is_reported_and_not_written(isolated_db, tmp_path, capsys):
    with repository.get_connection() as conn:
        source_id, roots = _source_and_tags(conn)
        _seed_feed(conn, source_id, roots, 3)
        conn.commit()
    _radar("only-tech", "technology", THEMES[0])
    out = tmp_path / "pack"

    _run("--out", str(out))
    assert "не хватает" in capsys.readouterr().out

    # 4 сигнала вместо 50 и радар одной категории — экзамен на таком наборе ничего не скажет.
    with pytest.raises(SystemExit):
        _run("--out", str(out), "--apply")
    assert not out.exists()


def test_pack_is_never_written_inside_the_repository_and_dates_are_checked(corpus):
    inside = Path(signal_card_pack.__file__).resolve().parents[1] / "exports" / "pack"

    with pytest.raises(SystemExit):
        _run("--out", str(inside), "--apply")
    assert not inside.exists()
    with pytest.raises(SystemExit):
        _run("--since", "август")


# --- Ручной отбор (16.10): владелец проверил 50 строк — выгрузка по списку номеров и ручные поправки.

def _ids_file(tmp_path, rows, name="ids.txt"):
    path = tmp_path / name
    path.write_text("# ручной отбор\n" + "\n".join(f"{item} {stratum}  # заголовок" for item, stratum in rows) + "\n",
                    encoding="utf-8")
    return path


def _listed(corpus):
    # Человек решает, полезна ли статья: одна с низким баллом взята полезной.
    noise = sorted(corpus["noise"])
    return ([(f"article:{i}", "useful") for i in sorted(corpus["useful"])[:23] + noise[:1]]
            + [(f"article:{i}", "noise") for i in noise[1:7]]
            + [(f"signal:{i}", "technology") for i in sorted(corpus["tech"])[:15]]
            + [(f"signal:{i}", "business") for i in sorted(corpus["business"])[:5]])


def test_ids_file_selects_exactly_the_listed_signals_with_their_strata(corpus, tmp_path):
    listed = _listed(corpus)
    out = tmp_path / "pack"

    _run("--out", str(out), "--apply", "--ids-file", str(_ids_file(tmp_path, listed)))

    manifest = _manifest(out)
    assert {(item["id"], item["stratum"]) for item in manifest["items"]} == set(listed)
    assert manifest["selection"]["ids_file"] == "ids.txt"


def test_ids_file_mistakes_are_refused(corpus, tmp_path):
    listed = _listed(corpus)
    tech, business = listed[30][0], listed[45][0]
    cases = {
        "unknown": (listed[:-1] + [("signal:999999", "business")], "нет среди кандидатов"),
        "not-a-candidate": ([(f"article:{sorted(corpus['not_feed'])[0]}", "useful")] + listed[1:], "нет среди кандидатов"),
        # Доля сигнала радара — его категория: технология, записанная бизнесом, — ошибка списка.
        "radar-stratum": ([(item, {tech: "business", business: "technology"}.get(item, stratum))
                           for item, stratum in listed], "доля не подходит"),
        "count": (listed[1:], "не те квоты"),
        "duplicate": (listed[:1] + listed[:1] + listed[2:], "дважды"),
    }
    for name, (rows, reason) in cases.items():
        out = tmp_path / name
        with pytest.raises(SystemExit) as refused:
            _run("--out", str(out), "--apply", "--ids-file", str(_ids_file(tmp_path, rows, f"{name}.txt")))
        assert reason in str(refused.value), (name, str(refused.value))
        assert not out.exists(), name


def test_overrides_fix_the_date_and_the_text_and_are_recorded(corpus, tmp_path):
    # Владелец 16.10: у строк радара дата публикации пуста (у находок поиска в базе — дата события по
    # версии модели), у части текстов в базе — меню сайта. Дату и текст берём с первоисточника вручную.
    listed = _listed(corpus)
    radar_id, feed_id = listed[30][0], listed[0][0]
    fixes = tmp_path / "fixes.json"
    fixes.write_text(json.dumps({
        radar_id: {"published": "2026-09-03", "text": "KUALA LUMPUR, 3 September 2026 — the release itself. " * 5,
                   "note": "дата и текст — со страницы источника"},
        feed_id: {"published": "2026-09-01"},
    }, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "pack"

    _run("--out", str(out), "--apply", "--ids-file", str(_ids_file(tmp_path, listed)), "--overrides", str(fixes))

    items = {item["id"]: item for item in _manifest(out)["items"]}
    radar = items[radar_id]
    assert radar["published"] == "2026-09-03" and radar["text"].startswith("KUALA LUMPUR")
    assert radar["text_sha256"] == hashlib.sha256(radar["text"].encode("utf-8")).hexdigest()
    assert radar["overrides"] == ["published", "text"] and "источника" in radar["override_note"]
    assert items[feed_id]["overrides"] == ["published"]
    assert not items[listed[1][0]]["overrides"]
    _, _, header, rows = _markup(out)
    row = next(row for row in rows if row[header.index("ID")].value == radar_id)
    assert row[header.index("Дата публикации")].value == "2026-09-03"
    assert row[header.index("Текст публикации")].value == radar["text"]

    for name, payload in {
        "not-in-sample": {"signal:999999": {"published": "2026-09-03"}},
        "bad-date": {radar_id: {"published": "03.09.2026"}},
        "unknown-field": {radar_id: {"category": "бизнес"}},
    }.items():
        bad = tmp_path / f"{name}.json"
        bad.write_text(json.dumps(payload), encoding="utf-8")
        with pytest.raises(SystemExit) as refused:
            _run("--out", str(tmp_path / name), "--apply", "--ids-file", str(_ids_file(tmp_path, listed)),
                 "--overrides", str(bad))
        assert "signal-card-expert-pack:" in str(refused.value), name  # отказ команды, а не разбор аргументов


def test_manifest_freezes_the_agent_inputs_and_the_table_names_the_publisher(corpus, tmp_path):
    # Экзамен #98 подаёт агенту входы из манифеста, а не из базы: поправки вручную видны только там.
    # Тип издателя нужен агенту #96 — компания-издатель участник своей новости, издание — нет.
    out = tmp_path / "pack"

    _run("--out", str(out), "--apply")

    items = _manifest(out)["items"]
    assert all({"title", "text", "published", "publisher", "publisher_type"} <= set(item) for item in items)
    feed = next(item for item in items if item["origin"] == "feed")
    radar = next(item for item in items if item["origin"] == "radar")
    assert (feed["publisher"], feed["publisher_type"]) == ("World Oil", "Media")
    assert radar["publisher_type"] == ""
    book, _, header, rows = _markup(out)
    by_id = {row[header.index("ID")].value: row for row in rows}
    assert by_id[feed["id"]][header.index("Издатель")].value == "World Oil"
    assert by_id[radar["id"]][header.index("Издатель")].value == "radar.example.com"
    guide = " ".join(str(row[0]) for row in book["Инструкция"].iter_rows(values_only=True) if row[0])
    assert "участником не считается" in guide and "пресс-релиз о своём событии" in guide

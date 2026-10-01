"""Эталон заказчика (xlsx ТОП-сигналов) — в память радара и в замер полноты.

Таблица здесь синтетическая: настоящая — у заказчика, репозиторий публичный."""

from __future__ import annotations

import openpyxl
import pytest

from oiltech_digest import cli, signal_discovery, signal_feedback, signal_reference
from oiltech_digest.db import repository

HEADER = ["№", "Направление", "НСР", "Сигнал / технология", "Краткая суть", "Почему важно для нефтесервиса",
          "Тип сигнала", "Стадия зрелости", "Приоритет", "Рекомендация", "Источник", "Ссылка"]


@pytest.fixture()
def reference_xlsx(tmp_path):
    book = openpyxl.Workbook()
    september = book.active
    september.title = "Сентябрь_2026"
    september.append(HEADER)
    september.append([1, "Целостность скважин", "", "Expro SafeWells V5 — единый контур контроля барьеров",
                      "Релиз 03.09", "Барьеры скважины в одном контуре", "", "Промышленное", 8.4,
                      "Проверить на фонде", "Expro", "https://www.expro.com/news/safewells-v5?utm_source=x"])
    september.append([2, "Транспорт / ПБ", "", "Hexagon VIS — вмешательство при угрозе столкновения",
                      "", "Перенос из горнодобычи", "", "", 9.4, "", "Hexagon", ""])
    september.append([None] * len(HEADER))  # пустая строка в конце листа
    notes = book.create_sheet("Заметки")
    notes.append(["Что-то", "без колонки сигнала"])
    august = book.create_sheet("Август_2026")
    august.append(HEADER)
    august.append([1, "ГРП", "", "Seismos — closed-loop ГРП по real-time акустике", "", "Управление ГРП в реальном времени",
                   "", "", 9.5, "ТОП", "JPT", "https://jpt.spe.org/seismos"])
    path = tmp_path / "top.xlsx"
    book.save(path)
    return path


def test_reads_signal_sheets_and_skips_others(reference_xlsx):
    rows = signal_reference.read_reference(reference_xlsx)

    assert [(row["sheet"], row["title"][:20]) for row in rows] == [
        ("Сентябрь_2026", "Expro SafeWells V5 —"), ("Сентябрь_2026", "Hexagon VIS — вмешат"),
        ("Август_2026", "Seismos — closed-loo"),
    ]
    assert signal_reference.read_reference(reference_xlsx, sheets=["Август_2026"])[0]["why"] == "Управление ГРП в реальном времени"


def test_import_is_a_dry_run_by_default_and_idempotent(isolated_db, reference_xlsx):
    rows = signal_reference.read_reference(reference_xlsx, sheets=["Август_2026"])

    assert signal_reference.import_reference(rows)["applied"] is False
    assert repository.list_signal_agent_memory(memory_type="signal_verdict") == []

    signal_reference.import_reference(rows, apply=True)
    signal_reference.import_reference(rows, apply=True)  # повтор — та же память

    memory = repository.list_signal_agent_memory(memory_type="signal_verdict")
    assert len(memory) == 1
    row = memory[0]
    assert (row["subject"], float(row["score"])) == ("strong_signal", 100.0)
    assert row["facts_json"]["origin"] == "customer_reference"
    assert row["facts_json"]["reason"] == "Управление ГРП в реальном времени Рекомендация заказчика: ТОП"


def test_imported_reference_reaches_the_research_prompt(isolated_db, reference_xlsx):
    signal_reference.import_reference(signal_reference.read_reference(reference_xlsx, sheets=["Август_2026"]), apply=True)

    block = signal_feedback.research_feedback_block("ГРП")

    assert "- Seismos — closed-loop ГРП по real-time акустике — Управление ГРП в реальном времени" in block


def test_recall_matches_by_link_or_by_names_and_reports_misses(reference_xlsx):
    rows = signal_reference.read_reference(reference_xlsx, sheets=["Сентябрь_2026"])
    cards = [
        # Ссылка та же, хвост utm и www — не помеха.
        {"id": 17, "title_ru": "Цифровая платформа барьеров", "urls": ["https://expro.com/news/safewells-v5"], "visible": True},
        {"id": 40, "title_ru": "Hexagon VIS предотвращает столкновения на руднике", "urls": [], "visible": False},
    ]

    result = signal_reference.recall(rows, cards)

    assert (result["total"], result["found"], result["found_visible"]) == (2, 2, 1)
    assert [item["signal_id"] for item in result["found_items"]] == [17, 40]


def test_recall_does_not_match_on_a_single_common_word(reference_xlsx):
    rows = signal_reference.read_reference(reference_xlsx, sheets=["Сентябрь_2026"])
    # «Expro» есть, «SafeWells» — нет: это другой продукт той же компании.
    cards = [{"id": 22, "title_ru": "Expro Velonix управляет скоростью поршня", "urls": [], "visible": True}]

    result = signal_reference.recall(rows, cards)

    assert result["found"] == 0
    assert len(result["missing_items"]) == 2


def test_recall_cli_reads_cards_from_the_radar(isolated_db, reference_xlsx, capsys):
    signal_id = repository.upsert_signal({"signal_key": "k", "title": "Expro SafeWells V5 launched", "theme": "Бурение"})
    repository.upsert_signal_evidence(signal_id, {"source_url": "https://example.com/x", "title": "t"})

    cli.main(["radar-recall", str(reference_xlsx), "--sheet", "Сентябрь_2026"])

    out = capsys.readouterr().out
    assert "найдено 1 из 2 (50%)" in out and f"+ #{signal_id}" in out and "- Hexagon VIS" in out


def test_import_cli_dry_run_then_apply(isolated_db, reference_xlsx, capsys):
    cli.main(["import-reference-signals", str(reference_xlsx), "--sheet", "Август_2026"])
    assert "applied=False" in capsys.readouterr().out
    assert repository.list_signal_agent_memory(memory_type="signal_verdict") == []

    cli.main(["import-reference-signals", str(reference_xlsx), "--sheet", "Август_2026", "--apply"])
    assert "applied=True" in capsys.readouterr().out
    assert len(repository.list_signal_agent_memory(memory_type="signal_verdict")) == 1


def test_research_mode_judges_every_verified_event_beyond_max_signals(monkeypatch):
    from oiltech_digest import signal_research
    from tests.test_signal_research import _Research, _event

    events = [_event(title=f"Событие {i}", source_url=f"https://example.com/{i}") for i in range(9)]
    monkeypatch.setattr(signal_research, "make_client", lambda offline: _Research(events))
    monkeypatch.setattr(signal_discovery, "_fetch_full_text",
                        lambda url, fallback_title="", **kwargs: {"ok": False, "error": "403", "title": "", "raw_text": ""})
    run_config = signal_discovery.SignalDiscoveryConfig(offline=True, dry_run=True, web_only=True,
                                                        search_mode="openai_web", max_signals=6)

    run = signal_discovery.run_discovery(run_config, {"topics": [{"name": "Бурение"}], "tags": []})

    # max_signals=6 отрезал бы три события до судьи.
    assert len(run["topics"][0]["candidates"]) == 9

"""Общее для тестов пакета эксперту «Отбора» (#129, шов 2 спеки #128): синтетика на изолированной базе —
статьи ленты, находки радара, материалы коллег, темы и критерии — и прежняя разметка карточки во временной
папке. Настоящий пакет уходит эксперту вне репозитория (он публичный). Тесты — test_otbor_pack.py
(выгрузка) и test_otbor_markup.py (проверка разметки)."""

from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

import openpyxl

from oiltech_digest import cli, otbor_pack_format as pack_format
from oiltech_digest.db import repository
from oiltech_digest.processing import seed

UNTIL = "2026-10-22"  # окно ленты — 30 дней до него: с 22.09
THEMES = {
    "Бурение": "Строительство скважин и буровое оборудование",
    "Добыча": "Механизированная добыча и внутрискважинное оборудование",
    "Геология": "Сейсморазведка и геофизика",
    "Экология": "Выбросы, отходы и промышленная безопасность",
}
# Значения машины и автора материала: ни одно не должно попасть в таблицу эксперта.
MACHINE = [
    "CARD_SUMMARY", "CARD_TITLE_RU", "RELEVANCE_REASON", "TAG_RATIONALE", "JUDGE_TITLE", "JUDGE_SUMMARY",
    "JUDGE_WHY", "JUDGE_APPLICATION", "JudgeCo", "RESEARCH_SUMMARY", "JUDGE_VERDICT", "FILTER_REASON",
    "MIXED_REASON", "SearchModelPublisher", "AUTHOR_COMMENT", "AuthorCo", "2026-01-03",
]
MODEL_DATE = "2026-01-03"  # дата события по версии модели поиска; у «старых» — и дата судьи
BLOCK_PAGE = "Доступ с вашего IP-адреса ограничен. Отключите VPN/прокси и обновите страницу. " * 3
CHINESE = "大庆油田采油八厂第四作业区407注采班中控室内，集输工端坐在操作台前，目光在数据大屏上快速扫过。" * 4


def at(day: str) -> datetime:
    return datetime.fromisoformat(f"{day}T12:00:00+00:00")


def run_pack(*args: str) -> None:
    cli.main(["otbor-expert-pack", *args])


def run_check(*args: str) -> None:
    cli.main(["otbor-markup-check", *args])


def manifest(out: Path) -> dict:
    return json.loads((out / pack_format.MANIFEST_NAME).read_text(encoding="utf-8"))


def markup(out: Path):
    book = openpyxl.load_workbook(out / pack_format.WORKBOOK_NAME)
    sheet = book[pack_format.MARKUP_SHEET]
    rows = list(sheet.iter_rows())
    return book, sheet, [cell.value for cell in rows[0]], rows[1:]


def _source(conn, name="World Oil", *, archived=False) -> int:
    return conn.execute(
        "INSERT INTO sources (name, source_type, url, enabled, parse_strategy, archived_at)"
        " VALUES (%s, 'Media', %s, TRUE, 'request', %s) RETURNING id",
        (name, f"https://{name.lower().replace(' ', '')}.example.org", at("2026-10-01") if archived else None),
    ).fetchone()[0]


def article(conn, source_id, key, day, *, relevant=True, text=None, tag_id=None, language="en") -> int:
    article_id = conn.execute(
        "INSERT INTO articles (source_id, title, url, raw_text, published_at, collected_at, language)"
        " VALUES (%s, %s, %s, %s, %s, %s, %s) RETURNING id",
        (source_id, f"Publisher title {key}", f"https://news.example.org/{key}",
         text or f"Publisher body of {key}. " * 40, at(day), at(day), language),
    ).fetchone()[0]
    if relevant is not None:
        conn.execute(
            "INSERT INTO article_cards (article_id, summary, title_ru, relevant, relevance_reason)"
            " VALUES (%s, 'CARD_SUMMARY', 'CARD_TITLE_RU', %s, 'RELEVANCE_REASON')", (article_id, relevant))
    if tag_id:
        conn.execute("INSERT INTO article_tags (article_id, tag_id, confidence, rationale)"
                     " VALUES (%s, %s, 0.9, 'TAG_RATIONALE')", (article_id, tag_id))
        conn.execute("INSERT INTO article_scores (article_id, total_score) VALUES (%s, 77)", (article_id,))
    return article_id


def radar(key, *, category="technology", stage=None, mixed=False, old=False, links=1, fetched=True,
          seen="2026-10-01", archived=False, corpus_article=None) -> int:
    """Находка поиска: ссылки — страницы, которые докачали (текст издателя), дата ссылки — модели поиска.
    corpus_article — вторая ссылка ведёт на статью корпуса (заголовок и издатель — её)."""
    signal_id = repository.upsert_signal({
        "signal_key": key, "title": f"JUDGE_TITLE {key}", "title_ru": f"JUDGE_TITLE {key}", "theme": "Бурение",
        "summary": "JUDGE_SUMMARY", "maturity": "watch", "score": 70, "companies": ["JudgeCo"],
    })
    for number in range(links):
        repository.upsert_signal_evidence(signal_id, {
            "source_url": f"https://www.radar{number}.example.com/{key}", "title": f"Page title {key} #{number}",
            "publisher": "SearchModelPublisher", "published_at": MODEL_DATE,
            "extracted_fact": f"Page text of {key} #{number}. " * 12 if fetched else "JUDGE_SUMMARY",
            "summary_ru": "RESEARCH_SUMMARY", "strength": 0.9 - number * 0.05,
            "article_id": corpus_article if number == 1 else None,
            "raw_payload": {"evidence_source": "openai_web_search", "full_text_fetched": fetched,
                            "research_summary": "RESEARCH_SUMMARY"},
        })
    with repository.get_connection() as conn:
        conn.execute(
            "UPDATE signals SET signal_category = %s, event_date = %s, why_interesting = 'JUDGE_WHY',"
            " oilfield_relevance = 'direct', oilfield_application = 'JUDGE_APPLICATION',"
            " verdict_reason = 'JUDGE_VERDICT', filter_stage = %s, filter_reason = 'FILTER_REASON',"
            " mixed_events = %s, mixed_events_reason = 'MIXED_REASON', first_seen_at = %s, archived_at = %s"
            " WHERE id = %s",
            (category, MODEL_DATE if old else "2026-09-25", stage, mixed, at(seen),
             at("2026-10-05") if archived else None, signal_id),
        )
        conn.commit()
    return signal_id


def material(conn, key, *, kind="event", url=None, day="2026-10-10", signal_id=None) -> int:
    return conn.execute(
        "INSERT INTO colleague_materials (kind, title, description, url, companies_json, verdict, verdict_comment,"
        " tg_chat_id, tg_message_id, tg_user_id, tg_author, signal_id, created_at)"
        " VALUES (%s, %s, %s, %s, '[\"AuthorCo\"]', 'strong_signal', 'AUTHOR_COMMENT', 1, %s, 7, 'Коллега', %s, %s)"
        " RETURNING id",
        (kind, f"Author title {key}", f"Author description of {key}: what happened and why it matters. " * 3,
         url, int(hashlib.sha256(key.encode()).hexdigest()[:8], 16), signal_id, at(day)),
    ).fetchone()[0]


def build_corpus(tmp_path: Path) -> dict:
    """Кандидатов каждой группы — ровно по квоте (исключения видны по составу), плюс то, что в пакет не идёт."""
    seed.seed_default_scoring_criteria()
    with repository.get_connection() as conn:
        for order, (name, description) in enumerate(THEMES.items()):
            conn.execute("INSERT INTO tags (name, description, sort_order) VALUES (%s, %s, %s)",
                         (name, description, order))
        conn.execute("INSERT INTO tags (name, description, enabled) VALUES ('Выключенная тема', 'старая', FALSE)")
        conn.execute("INSERT INTO tags (name, description) VALUES (%s, 'приёмник')",
                     (repository.SYSTEM_TAG_UNCLASSIFIED,))
        tag = conn.execute("SELECT id FROM tags WHERE name = 'Бурение'").fetchone()[0]
        source = _source(conn)
        archived_source = _source(conn, "Old Source", archived=True)
        days = [f"2026-10-{day:02d}" for day in range(1, 21)] + [f"2026-09-{day}" for day in range(23, 31)]
        rejected = [article(conn, source, f"rejected-{i}", days[i % len(days)], relevant=False) for i in range(50)]
        accepted = [article(conn, source, f"accepted-{i}", days[i % len(days)], tag_id=tag) for i in range(30)]
        card_feed = [article(conn, source, f"card-{i}", "2026-10-02", tag_id=tag) for i in range(4)]
        reprint_primary = rejected[0]
        excluded_feed = [
            article(conn, source, "rejected-early", "2026-09-20", relevant=False),  # до окна
            article(conn, source, "accepted-late", "2026-10-22", tag_id=tag),  # окно не включает --until
            article(conn, source, "undecided", "2026-10-03", relevant=None),  # гейт ещё не решал
            article(conn, archived_source, "archived-source", "2026-10-03", relevant=False),
            article(conn, source, "blocked", "2026-10-04", relevant=False, text=BLOCK_PAGE),
            article(conn, source, "chinese", "2026-10-04", relevant=False, text=CHINESE),
            article(conn, source, "short", "2026-10-04", relevant=False, text="Короткий анонс."),
        ]
        reprint = article(conn, source, "reprint", "2026-10-05", relevant=False)
        conn.execute("INSERT INTO article_reprints (article_id, primary_id) VALUES (%s, %s)", (reprint, reprint_primary))
        deleted = article(conn, source, "pending-deletion", "2026-10-05", relevant=False)
        conn.execute("UPDATE articles SET pending_deletion = TRUE WHERE id = %s", (deleted,))
        mismatch = article(conn, source, "mismatch", "2026-10-05", tag_id=tag)
        conn.execute("UPDATE articles SET full_text_status = 'mismatch' WHERE id = %s", (mismatch,))
        excluded_feed += [reprint, deleted, mismatch]
        # Статья корпуса — ссылка находки; собрана до окна ленты, её кандидатом не стать.
        corpus_link = article(conn, source, "corpus-link", "2026-09-10", tag_id=tag)
        conn.commit()
    filtered = {
        "judge": [radar(f"judge-{i}", stage="judge") for i in range(6)],
        "review": [radar(f"review-{i}", stage="review", category="business") for i in range(5)],
        "mixed": [radar(f"mixed-{i}", mixed=True, links=2) for i in range(4)],
        "old": [radar(f"old-{i}", old=True) for i in range(3)],
    }
    links = [radar(f"links-{i}", links=2 + i % 2, category=("technology", "business", "other")[i % 3])
             for i in range(9)]
    links.append(radar("links-many", links=9, corpus_article=corpus_link))
    excluded_radar = [
        # Других этапов отсева в пакете нет, сколько бы ссылок у находки ни было.
        radar("search-stage", stage="search", links=2), radar("judge-error", stage="judge_error", links=2),
        radar("archived-judge", stage="judge", archived=True), radar("not-fetched", stage="judge", fetched=False),
        radar("single-link", links=1), radar("late-judge", stage="judge", seen=UNTIL),
    ]
    merged = radar("merged-judge", stage="judge")
    card_radar = [radar("card-links", links=3), radar("card-judge", stage="judge")]
    with repository.get_connection() as conn:
        conn.execute("UPDATE signals SET merged_into_signal_id = %s WHERE id = %s", (links[0], merged))
        materials = [material(conn, "link", kind="link", url="https://www.colleague.example.net/news/1"),
                     material(conn, "event")]
        excluded_materials = [material(conn, "became-signal", signal_id=links[0]),
                              material(conn, "late", day=UNTIL)]
        conn.commit()
    card = write_card_pack(tmp_path / "card_pack", card_feed, card_radar)
    return {
        "rejected": set(rejected), "accepted": set(accepted), "card_feed": set(card_feed),
        "excluded_feed": set(excluded_feed), "corpus_link": corpus_link,
        "filtered": filtered, "links": links, "links_many": links[-1],
        "excluded_radar": set(excluded_radar + [merged]), "card_radar": card_radar,
        "materials": set(materials), "excluded_materials": set(excluded_materials), "card": card,
    }


def write_card_pack(folder: Path, feed_ids: list[int], radar_ids: list[int]) -> dict:
    """Прежний пакет карточки (#94): манифест с замороженными входами и заполненная экспертом таблица. Текст
    манифеста — не текст базы (у прежних строк были поправки с первоисточника)."""
    folder.mkdir(parents=True)
    relevance = ["прямая", "переносимая", "нет связи", "прямая", "прямая", "переносимая"]
    items = []
    for number, (origin, ident) in enumerate([("feed", one) for one in feed_ids] + [("radar", one) for one in radar_ids]):
        text = f"FROZEN text of card row {number}: what the expert saw. " * 20
        items.append({
            "id": f"article:{ident}" if origin == "feed" else f"signal:{ident}", "origin": origin,
            "stratum": "useful" if origin == "feed" else "technology",
            "article_id": ident if origin == "feed" else None, "signal_id": ident if origin == "radar" else None,
            "evidence_id": None, "url": f"https://frozen.example.org/{number}", "category": "бизнес",
            "theme": "Бурение", "title": f"FROZEN title {number}", "text": text, "published": "2026-09-30",
            "publisher": "Frozen Publisher", "publisher_type": "Media",
            "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(), "text_chars": len(text),
            "override_note": "дата с первоисточника" if number == 0 else None,
            "overrides": ["published"] if number == 0 else [],
        })
    manifest_path = folder / "signal_card_markup_manifest.json"
    manifest_path.write_text(json.dumps({"seed": 1, "items": items}, ensure_ascii=False), encoding="utf-8")
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Разметка"
    sheet.append(["ID", "Заголовок", "Текст публикации", "Тип события", "Связь с нефтесервисом",
                  "Комментарий эксперта"])
    for item, value in zip(items, relevance):
        sheet.append([item["id"], item["title"], item["text"], "контракт", value, "ок"])
    markup_path = folder / "signal_card_markup_completed.xlsx"
    book.save(markup_path)
    return {"manifest": manifest_path, "markup": markup_path, "items": items,
            "relevance": dict(zip((item["id"] for item in items), relevance))}


def pack_args(corpus: dict, out: Path, *extra: str) -> list[str]:
    return ["--out", str(out), "--until", UNTIL, "--card-manifest", str(corpus["card"]["manifest"]),
            "--card-markup", str(corpus["card"]["markup"]), *extra]

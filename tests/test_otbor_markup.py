"""Загрузка разметки эксперта для «Отбора» (#129, шов 2 спеки #128): проверка целостности и отчёт о нарушениях.

Шов — команда `otbor-markup-check` по пакету, который выгрузила `otbor-expert-pack` на изолированной базе
(синтетика — otbor_pack_support)."""

from __future__ import annotations

import shutil

import openpyxl
import pytest

from oiltech_digest import otbor_pack
from tests.otbor_pack_support import build_corpus, manifest, pack_args, run_check, run_pack

RUBRICS = otbor_pack.load_rubrics()["profiles"]
BUSINESS = [f"Бизнес: {one['name']}" for one in RUBRICS["business"]]
TECH = [f"Технология: {one['name']}" for one in RUBRICS["tech_radar"]]
TAKEN = {"Отрасль": "нефтесервис", "Связь с нефтесервисом": "прямая", "Брать или нет": "брать", "Тема": "Бурение",
         "Общая оценка": "Полезный сигнал", "Почему такой балл": "Крупный контракт на буровые услуги"}


@pytest.fixture()
def pack(isolated_db, tmp_path):
    corpus = build_corpus(tmp_path)
    out = tmp_path / "pack"
    run_pack(*pack_args(corpus, out, "--apply"))
    return {"corpus": corpus, "out": out, "items": manifest(out)["items"]}


def business(**overrides):
    return {**TAKEN, "Категория": "бизнес", **{column: 70 for column in BUSINESS}, **overrides}


def technology(**overrides):
    return {**TAKEN, "Категория": "технология", **{column: 85 for column in TECH[:-1]}, TECH[-1]: "нет даты",
            **overrides}


def rejected(**overrides):
    return {"Отрасль": "нефтегаз", "Связь с нефтесервисом": "нет связи", "Брать или нет": "не брать",
            "Причина отказа": "Нефтегаз без связи с нефтесервисом", "Категория": "бизнес",
            "Общая оценка": "Не релевантно", **overrides}


def by_the_rules(items):
    """Разметка по правилам: отказы, бизнес, технологии, «другое» без баллов, «разные события» у находки со ссылками."""
    answers = {}
    for number, item in enumerate(items):
        answers[item["id"]] = [rejected(), business(), technology(),
                               {**TAKEN, "Категория": "другое", "Почему такой балл": None}][number % 4]
    links = next(item["id"] for item in items if item["group"] == "radar_links")
    answers[links] = rejected(**{"Причина отказа": "Разные события в одной находке", "Отрасль": "нефтесервис"})
    # «Свежесть» — число, если дата события в тексте есть.
    first_tech = next(item_id for item_id, answer in answers.items() if answer.get("Категория") == "технология")
    answers[first_tech] = technology(**{TECH[-1]: 90})
    return answers


def fill(out, answers, *, name="otbor_markup_done.xlsx", edit=None):
    book = openpyxl.load_workbook(out / otbor_pack.WORKBOOK_NAME)
    sheet = book[otbor_pack.MARKUP_SHEET]
    header = [cell.value for cell in sheet[1]]
    for row in sheet.iter_rows(min_row=2):
        for column, value in answers.get(row[0].value, {}).items():
            row[header.index(column)].value = value
    if edit:
        edit(book, sheet, header)
    path = out / name
    book.save(path)
    return path


def test_a_markup_by_the_rules_passes(pack, capsys):
    path = fill(pack["out"], by_the_rules(pack["items"]))

    run_check("--markup", str(path))

    printed = capsys.readouterr().out
    assert "строк 116, нарушений 0" in printed
    assert "Утверждение: «да» 0, правки 0, без решения 18" in printed


def test_every_violation_is_reported_at_once_and_the_check_fails(pack, capsys):
    items = pack["items"]
    answers = by_the_rules(items)
    feed = [item["id"] for item in items if item["group"] in ("feed_rejected", "feed_accepted")]
    (outside, no_reason, reason_taken, mixed_single, score_missing, score_bad, both_profiles, no_why, no_theme,
     scored_rejected, fresh_empty, blank, other_scored, changed_text, dropped, doubled) = feed[:16]
    answers.update({
        outside: rejected(**{"Отрасль": "энергетика"}),
        no_reason: rejected(**{"Причина отказа": None}),
        reason_taken: business(**{"Причина отказа": "Старое событие"}),
        mixed_single: rejected(**{"Причина отказа": "Разные события в одной находке"}),
        score_missing: business(**{BUSINESS[2]: None}),
        score_bad: business(**{BUSINESS[0]: 120, BUSINESS[1]: "высокий"}),
        both_profiles: technology(**{BUSINESS[0]: 50}),
        no_why: business(**{"Почему такой балл": None}),
        no_theme: business(**{"Тема": None, "Общая оценка": None}),
        scored_rejected: rejected(**{BUSINESS[0]: 40}),
        fresh_empty: technology(**{TECH[-1]: None}),
        blank: {"Брать или нет": None, "Отрасль": None, "Связь с нефтесервисом": None, "Категория": None},
        other_scored: {**TAKEN, "Категория": "другое", TECH[0]: 60},
    })

    def edit(book, sheet, header):
        rows = {row[0].value: row for row in sheet.iter_rows(min_row=2)}
        rows[changed_text][header.index("Текст публикации")].value = "Текст, который поправил эксперт"
        sheet.delete_rows(rows[dropped][0].row)
        sheet.append(["article:999999", "Лишняя строка", "текст"])
        sheet.append([cell.value for cell in rows[doubled]])

    path = fill(pack["out"], answers, edit=edit)

    with pytest.raises(SystemExit) as failed:
        run_check("--markup", str(path))

    printed = capsys.readouterr().out
    expected = [
        f"{outside}: «Отрасль»: «энергетика» нет в списке",
        f"{no_reason}: «Причина отказа» не заполнена",
        f"{reason_taken}: у строки «брать» не должно быть причины отказа",
        f"{mixed_single}: «Разные события в одной находке» только у строки со ссылками находки",
        f"{score_missing}: «{BUSINESS[2]}» не заполнена",
        f"{score_bad}: «{BUSINESS[0]}»: «120» не целое число от 0 до 100",
        f"{score_bad}: «{BUSINESS[1]}»: «высокий» не целое число от 0 до 100",
        f"{both_profiles}: баллы не той категории: «{BUSINESS[0]}»",
        f"{no_why}: «Почему такой балл» не заполнена",
        f"{no_theme}: «Тема» не заполнена",
        f"{no_theme}: «Общая оценка» не заполнена",
        f"{scored_rejected}: баллы у строки «не брать»",
        f"{fresh_empty}: «{TECH[-1]}» не заполнена",
        f"{blank}: «Брать или нет» не заполнена",
        f"{blank}: «Отрасль» не заполнена",
        f"{blank}: «Связь с нефтесервисом» не заполнена",
        f"{blank}: «Категория» не заполнена",
        f"{other_scored}: баллы у категории «другое»",
        f"{changed_text}: текст публикации не тот, что в манифесте",
        f"{dropped}: строки нет в таблице",
        "article:999999: строки нет в манифесте",
        f"{doubled}: номер повторяется",
    ]
    for line in expected:
        assert line in printed, line
    assert f"нарушений {len(expected)}" in printed
    assert "нарушений" in str(failed.value)


def test_the_check_needs_the_pack_the_markup_came_from(pack, tmp_path):
    path = fill(pack["out"], {})
    alone = tmp_path / "alone"
    alone.mkdir()
    shutil.copy(path, alone / path.name)

    with pytest.raises(SystemExit) as refused:
        run_check("--markup", str(alone / path.name))
    assert "манифест" in str(refused.value)

    with pytest.raises(SystemExit) as refused:
        run_check("--markup", str(path), "--manifest", str(pack["corpus"]["card"]["manifest"]))
    assert "не манифест пакета «Отбора»" in str(refused.value)

    def without_sheet(book, sheet, header):
        del book[otbor_pack.MARKUP_SHEET]

    with pytest.raises(SystemExit) as refused:
        run_check("--markup", str(fill(pack["out"], {}, name="no_sheet.xlsx", edit=without_sheet)))
    assert f"нет листа «{otbor_pack.MARKUP_SHEET}»" in str(refused.value)

    def without_column(book, sheet, header):
        sheet.delete_cols(header.index("Категория") + 1)

    with pytest.raises(SystemExit) as refused:
        run_check("--markup", str(fill(pack["out"], {}, name="no_column.xlsx", edit=without_column)))
    assert "нет колонок Категория" in str(refused.value)


def test_expert_decisions_on_the_rules_are_summarized(pack, capsys):
    def decide(book, sheet, header):
        for name, changed in ((otbor_pack.APPROVAL_SHEET, {"Старое событие": "убрать"}),
                              (otbor_pack.RUBRICS_SHEET, {"Свежесть": "своя правка: от даты события до сегодня"})):
            approval = book[name]
            columns = [cell.value for cell in approval[1]]
            for row in approval.iter_rows(min_row=2):
                row[columns.index("Решение")].value = changed.get(row[1].value, "да")

    path = fill(pack["out"], by_the_rules(pack["items"]), edit=decide)

    run_check("--markup", str(path))

    printed = capsys.readouterr().out
    assert "Утверждение: «да» 16, правки 2, без решения 0" in printed
    assert "Старое событие: убрать" in printed
    assert "Свежесть: своя правка: от даты события до сегодня" in printed


def test_dictionaries_changed_since_the_pack_are_named(pack, capsys, monkeypatch):
    path = fill(pack["out"], by_the_rules(pack["items"]))
    current = otbor_pack.load_dictionaries()
    monkeypatch.setattr(otbor_pack, "load_dictionaries", lambda: {**current, "version": "2099-01-01.1"})

    run_check("--markup", str(path))

    printed = capsys.readouterr().out
    assert f"справочники пакета {current['version']}, сейчас 2099-01-01.1" in printed
